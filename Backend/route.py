import base64
import binascii
import mimetypes
import re
import traceback
import uuid
from datetime import datetime
from io import BytesIO
from pathlib import Path

from fastapi import (
    APIRouter,
    Body,
    Depends,
    File,
    Form,
    HTTPException,
    UploadFile,
    status,
)
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from docx import Document
from sqlalchemy import inspect, text
from sqlalchemy.exc import DataError, IntegrityError
from sqlalchemy.orm import Session

from database import get_db

from model import (
    ability_exists,
    ability_exists_for_other,
    ability_to_dict,
    advertising_media_exists,
    advertising_media_exists_for_other,
    advertising_media_to_dict,
    advertising_purpose_exists,
    advertising_purpose_exists_for_other,
    advertising_purpose_to_dict,
    announcement_type_exists,
    announcement_type_exists_for_other,
    announcement_type_to_dict,
    create_ability,
    create_advertising_media,
    create_advertising_purpose,
    create_announcement_type,
    create_hobby,
    create_ksa,
    create_ksa_category,
    create_position_grade,
    create_role_in_offense,
    create_language,
    create_location,
    create_meeting_location,
    create_meeting_type,
    create_office_level,
    create_office_type,
    create_requirement,
    create_user,
    delete_ability as hard_delete_ability,
    delete_advertising_media as hard_delete_advertising_media,
    delete_advertising_purpose as hard_delete_advertising_purpose,
    delete_announcement_type as hard_delete_announcement_type,
    delete_hobby as hard_delete_hobby,
    delete_ksa as hard_delete_ksa,
    delete_ksa_category as hard_delete_ksa_category,
    delete_position_grade as hard_delete_position_grade,
    delete_role_in_offense as hard_delete_role_in_offense,
    delete_language as hard_delete_language,
    delete_location as hard_delete_location,
    delete_meeting_location as hard_delete_meeting_location,
    delete_meeting_type as hard_delete_meeting_type,
    delete_office_level as hard_delete_office_level,
    delete_office_type as hard_delete_office_type,
    delete_requirement as hard_delete_requirement,
    get_abilities,
    get_ability_by_id,
    get_active_user_count,
    get_active_users,
    get_advertising_media_by_id,
    get_advertising_medias,
    get_advertising_purpose_by_id,
    get_advertising_purposes,
    get_announcement_type_by_id,
    get_announcement_types,
    get_hobbies,
    get_hobby_by_id,
    get_ksa_by_id,
    get_ksa_categories,
    get_ksa_category_by_id,
    get_position_grade_by_id,
    get_position_grades,
    get_role_in_offense_by_id,
    get_role_in_offenses,
    get_ksas,
    get_language_by_id,
    get_languages,
    get_location_by_id,
    get_locations,
    get_meeting_location_by_id,
    get_meeting_locations,
    get_meeting_type_by_id,
    get_meeting_types,
    get_office_level_by_id,
    get_office_levels,
    get_office_type_by_id,
    get_office_types,
    get_requirement_by_id,
    get_requirements,
    get_user_by_email,
    get_user_by_full_name,
    get_user_by_id,
    get_user_rights_map,
    hobby_exists,
    hobby_exists_for_other,
    hobby_to_dict,
    ksa_category_exists,
    ksa_category_exists_for_other,
    ksa_category_to_dict,
    position_grade_exists,
    position_grade_exists_for_other,
    position_grade_to_dict,
    role_in_offense_exists,
    role_in_offense_exists_for_other,
    role_in_offense_to_dict,
    ksa_exists,
    ksa_exists_for_other,
    ksa_to_dict,
    language_exists,
    language_exists_for_other,
    language_to_dict,
    location_exists,
    location_exists_for_other,
    location_to_dict,
    meeting_location_exists,
    meeting_location_exists_for_other,
    meeting_location_to_dict,
    meeting_type_exists,
    meeting_type_exists_for_other,
    meeting_type_to_dict,
    office_level_exists,
    office_level_exists_for_other,
    office_level_to_dict,
    office_type_exists,
    office_type_exists_for_other,
    office_type_to_dict,
    requirement_exists,
    requirement_exists_for_other,
    requirement_to_dict,
    set_user_rights_bulk,
    soft_delete_user,
    update_ability,
    update_advertising_media,
    update_advertising_purpose,
    update_announcement_type,
    update_hobby,
    update_ksa,
    update_ksa_category,
    update_position_grade,
    update_role_in_offense,
    update_language,
    update_location,
    update_meeting_location,
    update_meeting_type,
    update_office_level,
    update_office_type,
    update_requirement,
    update_user,

    # --- Salary masters ---
    caste_exists,
    caste_exists_for_other,
    caste_to_dict,
    create_caste,

    create_nature_of_work,
    create_religion,
    create_schedule_type,
    create_skin_tone,
    create_task_status,

    # Shift Timing
    create_shift_timing,

    delete_caste as hard_delete_caste,
    delete_nature_of_work as hard_delete_nature_of_work,
    delete_religion as hard_delete_religion,
    delete_schedule_type as hard_delete_schedule_type,
    delete_skin_tone as hard_delete_skin_tone,
    delete_task_status as hard_delete_task_status,

    # Shift Timing
    delete_shift_timing,

    get_caste_by_id,
    get_castes,
    get_nature_of_work_by_id,
    get_nature_of_works,
    get_religion_by_id,
    get_religions,
    get_schedule_type_by_id,
    get_schedule_types,
    get_skin_tone_by_id,
    get_skin_tones,
    get_task_status_by_id,
    get_task_statuses,

    # Shift Timing
    get_shift_timing_by_id,
    get_shift_timings,

    nature_of_work_exists,
    nature_of_work_exists_for_other,
    nature_of_work_to_dict,

    religion_exists,
    religion_exists_for_other,
    religion_to_dict,

    schedule_type_exists,
    schedule_type_exists_for_other,
    schedule_type_to_dict,

    skin_tone_exists,
    skin_tone_exists_for_other,
    skin_tone_to_dict,

    task_status_to_dict,

    # Shift Timing
    shift_timing_to_dict,

    update_caste,
    update_nature_of_work,
    update_religion,
    update_schedule_type,
    update_skin_tone,
    update_task_status,

    # Shift Timing
    update_shift_timing,

    user_right_to_dict,

    create_employee_relation,
    delete_employee_relation as hard_delete_employee_relation,
    employee_relation_exists,
    employee_relation_exists_for_other,
    employee_relation_to_dict,
    get_employee_relation_by_id,
    get_employee_relations,
    update_employee_relation,
    # --- ported from route_mine.py (salary / lookups / accounts) ---
    AcctAccount,
    LOOKUP_TABLE_CONFIG,
    SalaryStructure,
    create_employee_document,
    create_lookup_row,
    create_salary_employee,
    create_salary_structure,
    delete_employee_children,
    delete_employee_document,
    delete_salary_employee,
    delete_salary_structure,
    employee_contact_to_dict,
    employee_document_to_dict,
    find_employee_active_record,
    find_employee_period_overlap,
    get_employee_contacts,
    get_employee_document,
    get_employee_documents,
    get_lookup_row_by_pk,
    get_lookup_rows,
    get_salary_employee_by_id,
    get_salary_employees,
    get_salary_structure_by_id,
    get_salary_structure_date_bounds,
    get_salary_structures,
    get_salemp_relations,
    list_user_account_maps,
    lookup_label_exists,
    lookup_row_to_dict,
    replace_employee_contacts,
    replace_employee_relations,
    resolve_default_set_id,
    resolve_legacy_user_id,
    salary_employee_field_exists,
    salary_employee_to_dict,
    salary_structure_to_dict,
    salemp_relation_to_dict,
    search_salary_employees,
    update_salary_employee,
    update_salary_structure,
    upsert_user_account_map,
)

from schema import (
    AbilityCreateRequest,
    AbilityUpdateRequest,
    AdvertisingMediaCreateRequest,
    AdvertisingMediaUpdateRequest,
    AdvertisingPurposeCreateRequest,
    AdvertisingPurposeUpdateRequest,
    AnnouncementTypeCreateRequest,
    AnnouncementTypeUpdateRequest,
    CreateUserRequest,
    HobbyCreateRequest,
    HobbyUpdateRequest,
    KSACategoryCreateRequest,
    KSACategoryUpdateRequest,
    PositionGradeCreateRequest,
    PositionGradeUpdateRequest,
    RoleInOffenseCreateRequest,
    RoleInOffenseUpdateRequest,
    KSACreateRequest,
    KSAUpdateRequest,
    LanguageCreateRequest,
    LanguageUpdateRequest,
    LocationCreateRequest,
    LocationUpdateRequest,
    LoginRequest,
    MeetingLocationCreateRequest,
    MeetingLocationUpdateRequest,
    MeetingTypeCreateRequest,
    MeetingTypeUpdateRequest,
    OfficeLevelCreateRequest,
    OfficeLevelUpdateRequest,
    OfficeTypeCreateRequest,
    OfficeTypeUpdateRequest,
    RegisterRequest,
    RequirementCreateRequest,
    RequirementUpdateRequest,
    SetUserRightsRequest,
    UpdateUserRequest,
    CasteCreateRequest,
    CasteUpdateRequest,
    NatureOfWorkCreateRequest,
    NatureOfWorkUpdateRequest,
    ReligionCreateRequest,
    ReligionUpdateRequest,
    ScheduleTypeCreateRequest,
    ScheduleTypeUpdateRequest,
    SkinToneCreateRequest,
    SkinToneUpdateRequest,
    TaskStatusCreateRequest,
    TaskStatusUpdateRequest,
    UpdateUserStatusRequest,
    EmployeeRelationCreateRequest,
    EmployeeRelationUpdateRequest,
    ShiftTimingCreateRequest,
    ShiftTimingUpdateRequest,
    # --- ported from route_mine.py ---
    LookupCreateRequest,
    SalEmpContactListRequest,
    SalEmpRelationListRequest,
    SalEmployeeCreateRequest,
    SalEmployeeUpdateRequest,
    SalStructureCreateRequest,
    SalStructureUpdateRequest,
    UserAccountMapRequest,
)

from security import (
    create_access_token,
    get_current_user,
    hash_password,
    require_admin,
    require_permission,
    validate_password,
    verify_password,
)


# ==================================================
# ROUTER
# ==================================================

router = APIRouter()


# ==================================================
# EXPORT / PRINT HELPERS (Phase 3)
# ==================================================

def _export_to_excel(records, columns, sheet_title="Data"):
    """
    records: list of model objects
    columns: list of (header_display, attr_name)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    header_fill = PatternFill(
        start_color="366092",
        end_color="366092",
        fill_type="solid",
    )
    header_font = Font(bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center")

    for col_idx, (header, _) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for row_idx, record in enumerate(records, 2):
        for col_idx, (_, attr) in enumerate(columns, 1):
            value = getattr(record, attr, None)
            ws.cell(row=row_idx, column=col_idx, value=value)

    for col in ws.columns:
        max_len = 0
        letter = col[0].column_letter
        for cell in col:
            max_len = max(max_len, len(str(cell.value or "")))
        ws.column_dimensions[letter].width = min(max_len + 2, 50)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def _export_to_word(records, columns, title):
    """
    records: list of model objects
    columns: list of (header_display, attr_name)
    """
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    )

    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Light Grid Accent 1"

    header_cells = table.rows[0].cells
    for idx, (header, _) in enumerate(columns):
        header_cells[idx].text = header
        for paragraph in header_cells[idx].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True

    for record in records:
        row_cells = table.add_row().cells
        for col_idx, (_, attr) in enumerate(columns):
            value = getattr(record, attr, None)
            row_cells[col_idx].text = "" if value is None else str(value)

    stream = BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream


def _file_date():
    return datetime.now().strftime("%Y_%m_%d")




# ==================================================
# EXPORT / PRINT ENDPOINTS (Phase 3)
# ==================================================


@router.get("/abilities/export")
def export_abilities(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("ability", "export")
    ),
):
    records = get_abilities(db)
    columns = [("ID", "pkABId"), ("Abilities", "Abilities")]
    stream = _export_to_excel(records, columns, sheet_title="Ability")
    filename = f"export_Ability_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/abilities/print")
def print_abilities(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("ability", "print")
    ),
):
    records = get_abilities(db)
    columns = [("ID", "pkABId"), ("Abilities", "Abilities")]
    stream = _export_to_word(
        records, columns, title="Ability Report"
    )
    filename = f"print_Ability_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/locations/export")
def export_locations(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("location", "export")
    ),
):
    records = get_locations(db)
    columns = [("ID", "pkHLId"), ("Location", "Location")]
    stream = _export_to_excel(records, columns, sheet_title="Location")
    filename = f"export_Location_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/locations/print")
def print_locations(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("location", "print")
    ),
):
    records = get_locations(db)
    columns = [("ID", "pkHLId"), ("Location", "Location")]
    stream = _export_to_word(
        records, columns, title="Location Report"
    )
    filename = f"print_Location_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/hobbies/export")
def export_hobbies(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("hobby", "export")
    ),
):
    records = get_hobbies(db)
    columns = [("ID", "pkHId"), ("Hobby", "Hobby")]
    stream = _export_to_excel(records, columns, sheet_title="Hobby")
    filename = f"export_Hobby_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/hobbies/print")
def print_hobbies(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("hobby", "print")
    ),
):
    records = get_hobbies(db)
    columns = [("ID", "pkHId"), ("Hobby", "Hobby")]
    stream = _export_to_word(
        records, columns, title="Hobby Report"
    )
    filename = f"print_Hobby_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/announcement-types/export")
def export_announcement_types(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("announcement_type", "export")
    ),
):
    records = get_announcement_types(db)
    columns = [("ID", "pkATId"), ("Announcement Type", "AnnouncementType")]
    stream = _export_to_excel(records, columns, sheet_title="AnnouncementType")
    filename = f"export_AnnouncementType_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/announcement-types/print")
def print_announcement_types(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("announcement_type", "print")
    ),
):
    records = get_announcement_types(db)
    columns = [("ID", "pkATId"), ("Announcement Type", "AnnouncementType")]
    stream = _export_to_word(
        records, columns, title="AnnouncementType Report"
    )
    filename = f"print_AnnouncementType_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/office-levels/export")
def export_office_levels(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("office_level", "export")
    ),
):
    records = get_office_levels(db)
    columns = [("ID", "pkOLId"), ("Office Level", "OfficeLevel")]
    stream = _export_to_excel(records, columns, sheet_title="OfficeLevel")
    filename = f"export_OfficeLevel_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/office-levels/print")
def print_office_levels(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("office_level", "print")
    ),
):
    records = get_office_levels(db)
    columns = [("ID", "pkOLId"), ("Office Level", "OfficeLevel")]
    stream = _export_to_word(
        records, columns, title="OfficeLevel Report"
    )
    filename = f"print_OfficeLevel_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/meeting-types/export")
def export_meeting_types(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("meeting_type", "export")
    ),
):
    records = get_meeting_types(db)
    columns = [("ID", "pkMTId"), ("Meeting Type", "MeetingType")]
    stream = _export_to_excel(records, columns, sheet_title="MeetingType")
    filename = f"export_MeetingType_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/meeting-types/print")
def print_meeting_types(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("meeting_type", "print")
    ),
):
    records = get_meeting_types(db)
    columns = [("ID", "pkMTId"), ("Meeting Type", "MeetingType")]
    stream = _export_to_word(
        records, columns, title="MeetingType Report"
    )
    filename = f"print_MeetingType_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/languages/export")
def export_languages(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("language", "export")
    ),
):
    records = get_languages(db)
    columns = [("ID", "pkLId"), ("Language", "Language")]
    stream = _export_to_excel(records, columns, sheet_title="Language")
    filename = f"export_Language_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/languages/print")
def print_languages(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("language", "print")
    ),
):
    records = get_languages(db)
    columns = [("ID", "pkLId"), ("Language", "Language")]
    stream = _export_to_word(
        records, columns, title="Language Report"
    )
    filename = f"print_Language_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/requirements/export")
def export_requirements(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("requirement", "export")
    ),
):
    records = get_requirements(db)
    columns = [("ID", "pkRId"), ("Requirement", "Requirement")]
    stream = _export_to_excel(records, columns, sheet_title="Requirement")
    filename = f"export_Requirement_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/requirements/print")
def print_requirements(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("requirement", "print")
    ),
):
    records = get_requirements(db)
    columns = [("ID", "pkRId"), ("Requirement", "Requirement")]
    stream = _export_to_word(
        records, columns, title="Requirement Report"
    )
    filename = f"print_Requirement_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/advertising-medias/export")
def export_advertising_medias(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("advertising_media", "export")
    ),
):
    records = get_advertising_medias(db)
    columns = [("ID", "pkAMId"), ("Advertising Media", "AdvertisingMedia")]
    stream = _export_to_excel(records, columns, sheet_title="AdvertisingMedia")
    filename = f"export_AdvertisingMedia_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/advertising-medias/print")
def print_advertising_medias(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("advertising_media", "print")
    ),
):
    records = get_advertising_medias(db)
    columns = [("ID", "pkAMId"), ("Advertising Media", "AdvertisingMedia")]
    stream = _export_to_word(
        records, columns, title="AdvertisingMedia Report"
    )
    filename = f"print_AdvertisingMedia_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/advertising-purposes/export")
def export_advertising_purposes(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("advertising_purpose", "export")
    ),
):
    records = get_advertising_purposes(db)
    columns = [("ID", "pkAPId"), ("Advertising Purpose", "AdvertisingPurpose")]
    stream = _export_to_excel(records, columns, sheet_title="AdvertisingPurpose")
    filename = f"export_AdvertisingPurpose_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/advertising-purposes/print")
def print_advertising_purposes(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("advertising_purpose", "print")
    ),
):
    records = get_advertising_purposes(db)
    columns = [("ID", "pkAPId"), ("Advertising Purpose", "AdvertisingPurpose")]
    stream = _export_to_word(
        records, columns, title="AdvertisingPurpose Report"
    )
    filename = f"print_AdvertisingPurpose_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/office-types/export")
def export_office_types(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("office_type", "export")
    ),
):
    records = get_office_types(db)
    columns = [("ID", "pkOTId"), ("Office Type", "OfficeType")]
    stream = _export_to_excel(records, columns, sheet_title="OfficeType")
    filename = f"export_OfficeType_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/office-types/print")
def print_office_types(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("office_type", "print")
    ),
):
    records = get_office_types(db)
    columns = [("ID", "pkOTId"), ("Office Type", "OfficeType")]
    stream = _export_to_word(
        records, columns, title="OfficeType Report"
    )
    filename = f"print_OfficeType_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/meeting-locations/export")
def export_meeting_locations(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("meeting_location", "export")
    ),
):
    records = get_meeting_locations(db)
    columns = [("ID", "pkMLId"), ("Meeting Location", "MeetingLocation")]
    stream = _export_to_excel(records, columns, sheet_title="MeetingLocation")
    filename = f"export_MeetingLocation_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/meeting-locations/print")
def print_meeting_locations(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("meeting_location", "print")
    ),
):
    records = get_meeting_locations(db)
    columns = [("ID", "pkMLId"), ("Meeting Location", "MeetingLocation")]
    stream = _export_to_word(
        records, columns, title="MeetingLocation Report"
    )
    filename = f"print_MeetingLocation_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/ksas/export")
def export_ksas(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("ksa", "export")
    ),
):
    records = get_ksas(db)
    columns = [("ID", "pkKSAId"), ("KSA", "KSA")]
    stream = _export_to_excel(records, columns, sheet_title="KSA")
    filename = f"export_KSA_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/ksas/print")
def print_ksas(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("ksa", "print")
    ),
):
    records = get_ksas(db)
    columns = [("ID", "pkKSAId"), ("KSA", "KSA")]
    stream = _export_to_word(
        records, columns, title="KSA Report"
    )
    filename = f"print_KSA_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/ksa-categories/export")
def export_ksa_categories(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("ksa_category", "export")
    ),
):
    records = get_ksa_categories(db)
    columns = [("ID", "pkKSACId"), ("KSA Category", "KSACategory")]
    stream = _export_to_excel(records, columns, sheet_title="KSACategory")
    filename = f"export_KSACategory_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/ksa-categories/print")
def print_ksa_categories(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("ksa_category", "print")
    ),
):
    records = get_ksa_categories(db)
    columns = [("ID", "pkKSACId"), ("KSA Category", "KSACategory")]
    stream = _export_to_word(
        records, columns, title="KSACategory Report"
    )
    filename = f"print_KSACategory_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/position-grades/export")
def export_position_grades(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("position_grade", "export")
    ),
):
    records = get_position_grades(db)
    columns = [("ID", "pkPGId"), ("Position Grade", "PositionGrade"), ("Minimum Pay", "MinimumPay"), ("Maximum Pay", "MaximumPay")]
    stream = _export_to_excel(records, columns, sheet_title="PositionGrade")
    filename = f"export_PositionGrade_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/position-grades/print")
def print_position_grades(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("position_grade", "print")
    ),
):
    records = get_position_grades(db)
    columns = [("ID", "pkPGId"), ("Position Grade", "PositionGrade"), ("Minimum Pay", "MinimumPay"), ("Maximum Pay", "MaximumPay")]
    stream = _export_to_word(
        records, columns, title="PositionGrade Report"
    )
    filename = f"print_PositionGrade_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/role-in-offenses/export")
def export_role_in_offenses(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("role_in_offense", "export")
    ),
):
    records = get_role_in_offenses(db)
    columns = [("ID", "pkRIOId"), ("Role in Offense", "RoleInOffense"), ("Minimum Penalty", "MinimumPenalty"), ("Maximum Penalty", "MaximumPenalty")]
    stream = _export_to_excel(records, columns, sheet_title="RoleInOffense")
    filename = f"export_RoleInOffense_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/role-in-offenses/print")
def print_role_in_offenses(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("role_in_offense", "print")
    ),
):
    records = get_role_in_offenses(db)
    columns = [("ID", "pkRIOId"), ("Role in Offense", "RoleInOffense"), ("Minimum Penalty", "MinimumPenalty"), ("Maximum Penalty", "MaximumPenalty")]
    stream = _export_to_word(
        records, columns, title="RoleInOffense Report"
    )
    filename = f"print_RoleInOffense_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# ==================================================
# SALARY MASTERS — EXPORT / PRINT
# ==================================================

@router.get("/nature-of-works/export")
def export_nature_of_works(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_nature_of_work", "export")
    ),
):
    records = get_nature_of_works(db)
    columns = [("ID", "pkNWId"), ("Nature of Work", "NatureOfWork")]
    stream = _export_to_excel(records, columns, sheet_title="NatureOfWork")
    filename = f"export_NatureOfWork_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/nature-of-works/print")
def print_nature_of_works(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_nature_of_work", "print")
    ),
):
    records = get_nature_of_works(db)
    columns = [("ID", "pkNWId"), ("Nature of Work", "NatureOfWork")]
    stream = _export_to_word(
        records, columns, title="Nature of Work Report"
    )
    filename = f"print_NatureOfWork_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/schedule-types/export")
def export_schedule_types(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_schedule_type", "export")
    ),
):
    records = get_schedule_types(db)
    columns = [("ID", "pkSTId"), ("Schedule Type", "Type")]
    stream = _export_to_excel(records, columns, sheet_title="ScheduleType")
    filename = f"export_ScheduleType_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/schedule-types/print")
def print_schedule_types(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_schedule_type", "print")
    ),
):
    records = get_schedule_types(db)
    columns = [("ID", "pkSTId"), ("Schedule Type", "Type")]
    stream = _export_to_word(
        records, columns, title="Schedule Type Report"
    )
    filename = f"print_ScheduleType_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/religions/export")
def export_religions(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_religion", "export")
    ),
):
    records = get_religions(db)
    columns = [("ID", "pkRGId"), ("Religion", "Religion")]
    stream = _export_to_excel(records, columns, sheet_title="Religion")
    filename = f"export_Religion_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/religions/print")
def print_religions(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_religion", "print")
    ),
):
    records = get_religions(db)
    columns = [("ID", "pkRGId"), ("Religion", "Religion")]
    stream = _export_to_word(
        records, columns, title="Religion Report"
    )
    filename = f"print_Religion_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/castes/export")
def export_castes(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_caste", "export")
    ),
):
    records = get_castes(db)
    columns = [("ID", "pkCSId"), ("Caste", "Caste")]
    stream = _export_to_excel(records, columns, sheet_title="Caste")
    filename = f"export_Caste_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/castes/print")
def print_castes(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_caste", "print")
    ),
):
    records = get_castes(db)
    columns = [("ID", "pkCSId"), ("Caste", "Caste")]
    stream = _export_to_word(
        records, columns, title="Caste Report"
    )
    filename = f"print_Caste_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/skin-tones/export")
def export_skin_tones(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_skin_tone", "export")
    ),
):
    records = get_skin_tones(db)
    columns = [("ID", "pkSkinId"), ("Colour", "Colour")]
    stream = _export_to_excel(records, columns, sheet_title="SkinTone")
    filename = f"export_SkinTone_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/skin-tones/print")
def print_skin_tones(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_skin_tone", "print")
    ),
):
    records = get_skin_tones(db)
    columns = [("ID", "pkSkinId"), ("Colour", "Colour")]
    stream = _export_to_word(
        records, columns, title="Skin Tone Report"
    )
    filename = f"print_SkinTone_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/task-statuses/export")
def export_task_statuses(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_task_status", "export")
    ),
):
    records = get_task_statuses(db)
    columns = [("ID", "pkStaId"), ("Finish", "Finish"), ("Cancel", "Cancel")]
    stream = _export_to_excel(records, columns, sheet_title="TaskStatus")
    filename = f"export_TaskStatus_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/task-statuses/print")
def print_task_statuses(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_task_status", "print")
    ),
):
    records = get_task_statuses(db)
    columns = [("ID", "pkStaId"), ("Finish", "Finish"), ("Cancel", "Cancel")]
    stream = _export_to_word(
        records, columns, title="Task Status Report"
    )
    filename = f"print_TaskStatus_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )

@router.get("/employee-relations/export")
def export_employee_relations(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_employee_relation", "export")
    ),
):
    records = get_employee_relations(db)
    columns = [("ID", "pkRelId"), ("Relationship", "Relationship")]
    stream = _export_to_excel(records, columns, sheet_title="EmployeeRelation")
    filename = f"export_EmployeeRelation_{_file_date()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/employee-relations/print")
def print_employee_relations(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_employee_relation", "print")
    ),
):
    records = get_employee_relations(db)
    columns = [("ID", "pkRelId"), ("Relationship", "Relationship")]
    stream = _export_to_word(records, columns, title="Employee Relation Report")
    filename = f"print_EmployeeRelation_{_file_date()}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )

# ==================================================
# SALARY — SHIFT TIMING — EXPORT
# ==================================================

@router.get("/shift-timings/export")
def export_shift_timings(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_shift_timing", "export")
    ),
):
    records = get_shift_timings(db)

    columns = [
        ("ID", "pkSTId"),
        ("Shift", "Shift"),
        ("Start Work", "StartWork"),
        ("End Work", "EndWork"),
        ("Total Work", "TotalWork"),
        ("Start Break", "StartBreak"),
        ("End Break", "EndBreak"),
        ("Total Break", "TotalBreak"),
    ]

    stream = _export_to_excel(
        records,
        columns,
        sheet_title="ShiftTiming",
    )

    filename = f"export_ShiftTiming_{_file_date()}.xlsx"

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={filename}"
        },
    )


# ==================================================
# SALARY — SHIFT TIMING — PRINT
# ==================================================

@router.get("/shift-timings/print")
def print_shift_timings(
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_shift_timing", "print")
    ),
):
    records = get_shift_timings(db)

    columns = [
        ("ID", "pkSTId"),
        ("Shift", "Shift"),
        ("Start Work", "StartWork"),
        ("End Work", "EndWork"),
        ("Total Work", "TotalWork"),
        ("Start Break", "StartBreak"),
        ("End Break", "EndBreak"),
        ("Total Break", "TotalBreak"),
    ]

    stream = _export_to_word(
        records,
        columns,
        title="Shift Timing Report",
    )

    filename = f"print_ShiftTiming_{_file_date()}.docx"

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f"attachment; filename={filename}"
        },
    )


# ==================================================
# REGISTER
# ==================================================

@router.post(
    "/register",
    status_code=status.HTTP_201_CREATED,
)
def register(

    request: RegisterRequest,

    db: Session = Depends(get_db),
):

    # ==============================================
    # NORMALIZE EMAIL / FULL NAME
    # ==============================================

    email = request.email.strip().lower()

    full_name = request.full_name.strip()


    # ==============================================
    # CHECK EMAIL
    # ==============================================

    existing_email_user = get_user_by_email(
        db,
        email,
    )

    if existing_email_user:

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Email already registered",
        )


    # ==============================================
    # CHECK FULL NAME
    # (login is by full name, so it must be unique)
    # ==============================================

    existing_full_name_user = get_user_by_full_name(
        db,
        full_name,
    )

    if existing_full_name_user:

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="An account with this full name already exists",
        )


    # ==============================================
    # VALIDATE PASSWORD
    # ==============================================

    validate_password(
        request.password
    )


    # ==============================================
    # HASH PASSWORD
    # ==============================================

    password_hash = hash_password(
        request.password
    )


    # ==============================================
    # CREATE USER
    # ==============================================

    user = create_user(

        db=db,

        email=email,

        password_hash=password_hash,

        full_name=full_name,

        role="user",
    )


    # ==============================================
    # RESPONSE
    # ==============================================

    return {

        "message": "User registered successfully",

        "user": {

            "pkid": user.pkid,

            "email": user.email,

            "full_name": user.full_name,

            "role": user.role,

            "is_active": user.is_active,
        },
    }


# ==================================================
# LOGIN
# ==================================================

@router.post(
    "/login",
)
def login(

    request: LoginRequest,

    db: Session = Depends(get_db),
):

    # ==============================================
    # NORMALIZE FULL NAME
    # ==============================================

    full_name = request.full_name.strip()


    # ==============================================
    # GET USER
    # ==============================================

    user = get_user_by_full_name(
        db,
        full_name,
    )

    if not user:

        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid full name or password",
        )


    # ==============================================
    # DELETED USER CHECK
    # ==============================================

    if user.deleted_at is not None:

        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid full name or password",
        )


    # ==============================================
    # ACTIVE USER CHECK
    # ==============================================

    if not user.is_active:

        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Your account has been deactivated",
        )


    # ==============================================
    # VERIFY PASSWORD
    # ==============================================

    password_valid = verify_password(

        request.password,

        user.password_hash,
    )

    if not password_valid:

        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid full name or password",
        )


    # ==============================================
    # CREATE TOKEN
    # ==============================================

    access_token = create_access_token(
        user.pkid
    )


    # ==============================================
    # RESPONSE
    # ==============================================

    return {

        "message": "Login successful",

        "access_token": access_token,

        "token_type": "bearer",

        "user": {

            "pkid": user.pkid,

            "email": user.email,

            "full_name": user.full_name,

            "role": user.role,

            "is_active": user.is_active,
        },
    }


# ==================================================
# GET CURRENT USER
# ==================================================

@router.get(
    "/me",
)
def get_me(

    current_user: dict = Depends(
        get_current_user
    ),
):

    return {

        "user": current_user,
    }


# ==================================================
# GET ALL USERS
# ADMIN ONLY
# ==================================================

@router.get(
    "/users",
)
def get_users(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_admin
    ),
):

    users = get_active_users(
        db
    )

    return {

        "total": len(users),

        "users": [

            {

                "pkid": user.pkid,

                "email": user.email,

                "full_name": user.full_name,

                "role": user.role,

                "is_active": user.is_active,

                "updated_at": user.updated_at,
            }

            for user in users
        ],
    }


# ==================================================
# GET USER COUNT
# ADMIN ONLY
# MUST BE REGISTERED BEFORE /users/{user_id}
# OR "count" WILL BE PARSED AS user_id
# AND FAIL INT VALIDATION (422)
# ==================================================

@router.get(
    "/users/count",
)
def get_user_count(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_admin
    ),
):

    count = get_active_user_count(
        db
    )

    return {

        "count": count,
    }


# ==================================================
# GET USER BY ID
# ADMIN ONLY
# ==================================================

@router.get(
    "/users/{user_id}",
)
def get_user(

    user_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_admin
    ),
):

    user = get_user_by_id(
        db,
        user_id,
    )

    if not user:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="User not found",
        )

    return {

        "user": {

            "pkid": user.pkid,

            "email": user.email,

            "full_name": user.full_name,

            "role": user.role,

            "is_active": user.is_active,

            "updated_at": user.updated_at,
        },
    }


# ==================================================
# CREATE USER
# ADMIN ONLY
# ==================================================

@router.post(
    "/users",
    status_code=status.HTTP_201_CREATED,
)
def create_new_user(

    request: CreateUserRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_admin
    ),
):

    # ==============================================
    # NORMALIZE EMAIL / FULL NAME
    # ==============================================

    email = request.email.strip().lower()

    full_name = request.full_name.strip()


    # ==============================================
    # CHECK EMAIL
    # ==============================================

    existing_email_user = get_user_by_email(
        db,
        email,
    )

    if existing_email_user:

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Email already registered",
        )


    # ==============================================
    # CHECK FULL NAME
    # (login is by full name, so it must be unique)
    # ==============================================

    existing_full_name_user = get_user_by_full_name(
        db,
        full_name,
    )

    if existing_full_name_user:

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="An account with this full name already exists",
        )


    # ==============================================
    # VALIDATE PASSWORD
    # ==============================================

    validate_password(
        request.password
    )


    # ==============================================
    # HASH PASSWORD
    # ==============================================

    password_hash = hash_password(
        request.password
    )


    # ==============================================
    # CREATE USER
    # ==============================================

    user = create_user(

        db=db,

        email=email,

        password_hash=password_hash,

        full_name=full_name,

        role=request.role.strip().lower(),
    )


    return {

        "message": "User created successfully",

        "user": {

            "pkid": user.pkid,

            "email": user.email,

            "full_name": user.full_name,

            "role": user.role,

            "is_active": user.is_active,

            "updated_at": user.updated_at,
        },
    }


# ==================================================
# UPDATE USER
# ADMIN ONLY
# ==================================================

@router.put(
    "/users/{user_id}",
)
def update_existing_user(

    user_id: int,

    request: UpdateUserRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_admin
    ),
):

    # ==============================================
    # GET USER
    # ==============================================

    user = get_user_by_id(
        db,
        user_id,
    )

    if not user:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="User not found",
        )


    # ==============================================
    # UPDATE DATA
    # ==============================================

    update_data = {}


    # ==============================================
    # FULL NAME
    # (login is by full name, so it must stay unique)
    # ==============================================

    if request.full_name is not None:

        full_name = request.full_name.strip()

        existing_full_name_user = get_user_by_full_name(
            db,
            full_name,
        )

        if (

            existing_full_name_user

            and existing_full_name_user.pkid != user.pkid

        ):

            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="An account with this full name already exists",
            )

        update_data["full_name"] = full_name


    # ==============================================
    # EMAIL
    # ==============================================

    if request.email is not None:

        email = request.email.strip().lower()

        existing_user = get_user_by_email(
            db,
            email,
        )

        if (

            existing_user

            and existing_user.pkid != user.pkid

        ):

            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Email already registered",
            )

        update_data["email"] = email


    # ==============================================
    # ROLE
    # ==============================================

    if request.role is not None:

        update_data["role"] = (
            request.role.strip().lower()
        )


    # ==============================================
    # PASSWORD
    # ==============================================

    if request.password is not None:

        validate_password(
            request.password
        )

        update_data["password_hash"] = (
            hash_password(
                request.password
            )
        )


    # ==============================================
    # UPDATE USER
    # ==============================================

    user = update_user(

        db,

        user,

        **update_data,
    )


    return {

        "message": "User updated successfully",

        "user": {

            "pkid": user.pkid,

            "email": user.email,

            "full_name": user.full_name,

            "role": user.role,

            "is_active": user.is_active,

            "updated_at": user.updated_at,
        },
    }


# ==================================================
# ACTIVATE / DEACTIVATE USER
# ADMIN ONLY
# ==================================================

@router.patch(
    "/users/{user_id}/status",
)
def update_user_status(

    user_id: int,

    request: UpdateUserStatusRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_admin
    ),
):

    user = get_user_by_id(
        db,
        user_id,
    )

    if not user:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="User not found",
        )


    # ==============================================
    # UPDATE STATUS
    # ==============================================

    user.is_active = request.is_active

    user = update_user(
        db,
        user,
    )


    return {

        "message": (
            "User activated successfully"
            if request.is_active
            else "User deactivated successfully"
        ),

        "user": {

            "pkid": user.pkid,

            "email": user.email,

            "full_name": user.full_name,

            "role": user.role,

            "is_active": user.is_active,
        },
    }


# ==================================================
# SOFT DELETE USER
# ADMIN ONLY
# ==================================================

@router.delete(
    "/users/{user_id}",
)
def delete_user(

    user_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_admin
    ),
):

    user = get_user_by_id(
        db,
        user_id,
    )

    if not user:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="User not found",
        )


    # ==============================================
    # PREVENT ADMIN FROM DELETING SELF
    # ==============================================

    if user.pkid == current_user["pkid"]:

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="You cannot delete your own account",
        )


    # ==============================================
    # SOFT DELETE
    # ==============================================

    soft_delete_user(
        db,
        user,
    )


    return {

        "message": "User deleted successfully",
    }


# ==================================================
# GET ABILITIES
# ==================================================

@router.get(
    "/abilities",
)
def list_abilities(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    abilities = get_abilities(
        db
    )

    return {

        "total": len(abilities),

        "abilities": [

            ability_to_dict(ability)

            for ability in abilities
        ],
    }


# ==================================================
# GET ABILITY BY ID
# ==================================================

@router.get(
    "/abilities/{ability_id}",
)
def get_ability(

    ability_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    ability = get_ability_by_id(
        db,
        ability_id,
    )

    if not ability:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Ability not found",
        )

    return {

        "ability": ability_to_dict(ability),
    }


# ==================================================
# CREATE ABILITY
# REQUIRES: ability / add
# ==================================================

@router.post(
    "/abilities",
    status_code=status.HTTP_201_CREATED,
)
def create_new_ability(

    request: AbilityCreateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("ability", "add")
    ),
):

    if ability_exists(
        db,
        request.abilities,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Ability already exists",
        )

    ability = create_ability(
        db,
        request.abilities,
        fk_user_id=resolve_legacy_user_id(db, current_user["pkid"]),
    )

    return {

        "message": "Ability created successfully",

        "ability": ability_to_dict(ability),
    }


# ==================================================
# UPDATE ABILITY
# REQUIRES: ability / edit
# ==================================================

@router.put(
    "/abilities/{ability_id}",
)
def update_existing_ability(

    ability_id: int,

    request: AbilityUpdateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("ability", "edit")
    ),
):

    existing = get_ability_by_id(
        db,
        ability_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Ability not found",
        )

    if ability_exists_for_other(
        db,
        request.abilities,
        ability_id,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Ability already exists",
        )

    ability = update_ability(
        db,
        ability_id,
        request.abilities,
    )

    return {

        "message": "Ability updated successfully",

        "ability": ability_to_dict(ability),
    }


# ==================================================
# DELETE ABILITY
# REQUIRES: ability / delete
# HARD DELETE (Phase 2)
# ==================================================

@router.delete(
    "/abilities/{ability_id}",
)
def delete_ability(

    ability_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("ability", "delete")
    ),
):

    existing = get_ability_by_id(
        db,
        ability_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Ability not found",
        )

    hard_delete_ability(
        db,
        ability_id,
    )

    return {

        "message": "Ability deleted successfully",
    }


# ==================================================
# GET LOCATIONS
# ==================================================

@router.get(
    "/locations",
)
def list_locations(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    locations = get_locations(
        db
    )

    return {

        "total": len(locations),

        "locations": [

            location_to_dict(location)

            for location in locations
        ],
    }


# ==================================================
# GET LOCATION BY ID
# ==================================================

@router.get(
    "/locations/{location_id}",
)
def get_location(

    location_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    location = get_location_by_id(
        db,
        location_id,
    )

    if not location:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Location not found",
        )

    return {

        "location": location_to_dict(location),
    }


# ==================================================
# CREATE LOCATION
# REQUIRES: location / add
# ==================================================

@router.post(
    "/locations",
    status_code=status.HTTP_201_CREATED,
)
def create_new_location(

    request: LocationCreateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("location", "add")
    ),
):

    if location_exists(
        db,
        request.location,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Location already exists",
        )

    location = create_location(
        db,
        request.location,
        fk_user_id=resolve_legacy_user_id(db, current_user["pkid"]),
    )

    return {

        "message": "Location created successfully",

        "location": location_to_dict(location),
    }


# ==================================================
# UPDATE LOCATION
# REQUIRES: location / edit
# ==================================================

@router.put(
    "/locations/{location_id}",
)
def update_existing_location(

    location_id: int,

    request: LocationUpdateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("location", "edit")
    ),
):

    existing = get_location_by_id(
        db,
        location_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Location not found",
        )

    if location_exists_for_other(
        db,
        request.location,
        location_id,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Location already exists",
        )

    location = update_location(
        db,
        location_id,
        request.location,
    )

    return {

        "message": "Location updated successfully",

        "location": location_to_dict(location),
    }


# ==================================================
# DELETE LOCATION
# REQUIRES: location / delete
# HARD DELETE (Phase 2)
# ==================================================

@router.delete(
    "/locations/{location_id}",
)
def delete_location(

    location_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("location", "delete")
    ),
):

    existing = get_location_by_id(
        db,
        location_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Location not found",
        )

    hard_delete_location(
        db,
        location_id,
    )

    return {

        "message": "Location deleted successfully",
    }


# ==================================================
# GET HOBBIES
# ==================================================

@router.get(
    "/hobbies",
)
def list_hobbies(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    hobbies = get_hobbies(
        db
    )

    return {

        "total": len(hobbies),

        "hobbies": [

            hobby_to_dict(hobby)

            for hobby in hobbies
        ],
    }


# ==================================================
# GET HOBBY BY ID
# ==================================================

@router.get(
    "/hobbies/{hobby_id}",
)
def get_hobby(

    hobby_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    hobby = get_hobby_by_id(
        db,
        hobby_id,
    )

    if not hobby:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Hobby not found",
        )

    return {

        "hobby": hobby_to_dict(hobby),
    }


# ==================================================
# CREATE HOBBY
# REQUIRES: hobby / add
# ==================================================

@router.post(
    "/hobbies",
    status_code=status.HTTP_201_CREATED,
)
def create_new_hobby(

    request: HobbyCreateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("hobby", "add")
    ),
):

    if hobby_exists(
        db,
        request.hobby,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Hobby already exists",
        )

    hobby = create_hobby(
        db,
        request.hobby,
        fk_user_id=resolve_legacy_user_id(db, current_user["pkid"]),
    )

    return {

        "message": "Hobby created successfully",

        "hobby": hobby_to_dict(hobby),
    }


# ==================================================
# UPDATE HOBBY
# REQUIRES: hobby / edit
# ==================================================

@router.put(
    "/hobbies/{hobby_id}",
)
def update_existing_hobby(

    hobby_id: int,

    request: HobbyUpdateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("hobby", "edit")
    ),
):

    existing = get_hobby_by_id(
        db,
        hobby_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Hobby not found",
        )

    if hobby_exists_for_other(
        db,
        request.hobby,
        hobby_id,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Hobby already exists",
        )

    hobby = update_hobby(
        db,
        hobby_id,
        request.hobby,
    )

    return {

        "message": "Hobby updated successfully",

        "hobby": hobby_to_dict(hobby),
    }


# ==================================================
# DELETE HOBBY
# REQUIRES: hobby / delete
# HARD DELETE (Phase 2)
# ==================================================

@router.delete(
    "/hobbies/{hobby_id}",
)
def delete_hobby(

    hobby_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("hobby", "delete")
    ),
):

    existing = get_hobby_by_id(
        db,
        hobby_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Hobby not found",
        )

    hard_delete_hobby(
        db,
        hobby_id,
    )

    return {

        "message": "Hobby deleted successfully",
    }


# ==================================================
# GET ANNOUNCEMENT TYPES
# ==================================================

@router.get(
    "/announcement-types",
)
def list_announcement_types(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    announcement_types = get_announcement_types(
        db
    )

    return {

        "total": len(announcement_types),

        "announcement_types": [

            announcement_type_to_dict(announcement_type)

            for announcement_type in announcement_types
        ],
    }


# ==================================================
# GET ANNOUNCEMENT TYPE BY ID
# ==================================================

@router.get(
    "/announcement-types/{announcement_type_id}",
)
def get_announcement_type(

    announcement_type_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    announcement_type = get_announcement_type_by_id(
        db,
        announcement_type_id,
    )

    if not announcement_type:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Announcement type not found",
        )

    return {

        "announcement_type": announcement_type_to_dict(
            announcement_type
        ),
    }


# ==================================================
# CREATE ANNOUNCEMENT TYPE
# REQUIRES: announcement_type / add
# ==================================================

@router.post(
    "/announcement-types",
    status_code=status.HTTP_201_CREATED,
)
def create_new_announcement_type(

    request: AnnouncementTypeCreateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("announcement_type", "add")
    ),
):

    if announcement_type_exists(
        db,
        request.announcement_type,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Announcement type already exists",
        )

    announcement_type = create_announcement_type(
        db,
        request.announcement_type,
    )

    return {

        "message": "Announcement type created successfully",

        "announcement_type": announcement_type_to_dict(
            announcement_type
        ),
    }


# ==================================================
# UPDATE ANNOUNCEMENT TYPE
# REQUIRES: announcement_type / edit
# ==================================================

@router.put(
    "/announcement-types/{announcement_type_id}",
)
def update_existing_announcement_type(

    announcement_type_id: int,

    request: AnnouncementTypeUpdateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("announcement_type", "edit")
    ),
):

    existing = get_announcement_type_by_id(
        db,
        announcement_type_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Announcement type not found",
        )

    if announcement_type_exists_for_other(
        db,
        request.announcement_type,
        announcement_type_id,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Announcement type already exists",
        )

    announcement_type = update_announcement_type(
        db,
        announcement_type_id,
        request.announcement_type,
    )

    return {

        "message": "Announcement type updated successfully",

        "announcement_type": announcement_type_to_dict(
            announcement_type
        ),
    }


# ==================================================
# DELETE ANNOUNCEMENT TYPE
# REQUIRES: announcement_type / delete
# HARD DELETE (Phase 2)
# ==================================================

@router.delete(
    "/announcement-types/{announcement_type_id}",
)
def delete_announcement_type(

    announcement_type_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("announcement_type", "delete")
    ),
):

    existing = get_announcement_type_by_id(
        db,
        announcement_type_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Announcement type not found",
        )

    hard_delete_announcement_type(
        db,
        announcement_type_id,
    )

    return {

        "message": "Announcement type deleted successfully",
    }


# ==================================================
# GET OFFICE LEVELS
# ==================================================

@router.get(
    "/office-levels",
)
def list_office_levels(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    office_levels = get_office_levels(
        db
    )

    return {

        "total": len(office_levels),

        "office_levels": [

            office_level_to_dict(office_level)

            for office_level in office_levels
        ],
    }


# ==================================================
# GET OFFICE LEVEL BY ID
# ==================================================

@router.get(
    "/office-levels/{office_level_id}",
)
def get_office_level(

    office_level_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    office_level = get_office_level_by_id(
        db,
        office_level_id,
    )

    if not office_level:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Office level not found",
        )

    return {

        "office_level": office_level_to_dict(office_level),
    }


# ==================================================
# CREATE OFFICE LEVEL
# REQUIRES: office_level / add
# ==================================================

@router.post(
    "/office-levels",
    status_code=status.HTTP_201_CREATED,
)
def create_new_office_level(

    request: OfficeLevelCreateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("office_level", "add")
    ),
):

    if office_level_exists(
        db,
        request.office_level,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Office level already exists",
        )

    office_level = create_office_level(
        db,
        request.office_level,
    )

    return {

        "message": "Office level created successfully",

        "office_level": office_level_to_dict(office_level),
    }


# ==================================================
# UPDATE OFFICE LEVEL
# REQUIRES: office_level / edit
# ==================================================

@router.put(
    "/office-levels/{office_level_id}",
)
def update_existing_office_level(

    office_level_id: int,

    request: OfficeLevelUpdateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("office_level", "edit")
    ),
):

    existing = get_office_level_by_id(
        db,
        office_level_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Office level not found",
        )

    if office_level_exists_for_other(
        db,
        request.office_level,
        office_level_id,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Office level already exists",
        )

    office_level = update_office_level(
        db,
        office_level_id,
        request.office_level,
    )

    return {

        "message": "Office level updated successfully",

        "office_level": office_level_to_dict(office_level),
    }


# ==================================================
# DELETE OFFICE LEVEL
# REQUIRES: office_level / delete
# HARD DELETE (Phase 2)
# ==================================================

@router.delete(
    "/office-levels/{office_level_id}",
)
def delete_office_level(

    office_level_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("office_level", "delete")
    ),
):

    existing = get_office_level_by_id(
        db,
        office_level_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Office level not found",
        )

    hard_delete_office_level(
        db,
        office_level_id,
    )

    return {

        "message": "Office level deleted successfully",
    }


# ==================================================
# GET MEETING TYPES
# ==================================================

@router.get(
    "/meeting-types",
)
def list_meeting_types(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    meeting_types = get_meeting_types(
        db
    )

    return {

        "total": len(meeting_types),

        "meeting_types": [

            meeting_type_to_dict(meeting_type)

            for meeting_type in meeting_types
        ],
    }


# ==================================================
# GET MEETING TYPE BY ID
# ==================================================

@router.get(
    "/meeting-types/{meeting_type_id}",
)
def get_meeting_type(

    meeting_type_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    meeting_type = get_meeting_type_by_id(
        db,
        meeting_type_id,
    )

    if not meeting_type:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Meeting type not found",
        )

    return {

        "meeting_type": meeting_type_to_dict(meeting_type),
    }


# ==================================================
# CREATE MEETING TYPE
# REQUIRES: meeting_type / add
# ==================================================

@router.post(
    "/meeting-types",
    status_code=status.HTTP_201_CREATED,
)
def create_new_meeting_type(

    request: MeetingTypeCreateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("meeting_type", "add")
    ),
):

    if meeting_type_exists(
        db,
        request.meeting_type,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Meeting type already exists",
        )

    meeting_type = create_meeting_type(
        db,
        request.meeting_type,
        fk_user_id=resolve_legacy_user_id(db, current_user["pkid"]),
    )

    return {

        "message": "Meeting type created successfully",

        "meeting_type": meeting_type_to_dict(meeting_type),
    }


# ==================================================
# UPDATE MEETING TYPE
# REQUIRES: meeting_type / edit
# ==================================================

@router.put(
    "/meeting-types/{meeting_type_id}",
)
def update_existing_meeting_type(

    meeting_type_id: int,

    request: MeetingTypeUpdateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("meeting_type", "edit")
    ),
):

    existing = get_meeting_type_by_id(
        db,
        meeting_type_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Meeting type not found",
        )

    if meeting_type_exists_for_other(
        db,
        request.meeting_type,
        meeting_type_id,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Meeting type already exists",
        )

    meeting_type = update_meeting_type(
        db,
        meeting_type_id,
        request.meeting_type,
    )

    return {

        "message": "Meeting type updated successfully",

        "meeting_type": meeting_type_to_dict(meeting_type),
    }


# ==================================================
# DELETE MEETING TYPE
# REQUIRES: meeting_type / delete
# HARD DELETE (Phase 2)
# ==================================================

@router.delete(
    "/meeting-types/{meeting_type_id}",
)
def delete_meeting_type(

    meeting_type_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("meeting_type", "delete")
    ),
):

    existing = get_meeting_type_by_id(
        db,
        meeting_type_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Meeting type not found",
        )

    hard_delete_meeting_type(
        db,
        meeting_type_id,
    )

    return {

        "message": "Meeting type deleted successfully",
    }


# ==================================================
# GET LANGUAGES
# ==================================================

@router.get(
    "/languages",
)
def list_languages(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    languages = get_languages(
        db
    )

    return {

        "total": len(languages),

        "languages": [

            language_to_dict(language)

            for language in languages
        ],
    }


# ==================================================
# GET LANGUAGE BY ID
# ==================================================

@router.get(
    "/languages/{language_id}",
)
def get_language(

    language_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    language = get_language_by_id(
        db,
        language_id,
    )

    if not language:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Language not found",
        )

    return {

        "language": language_to_dict(language),
    }


# ==================================================
# CREATE LANGUAGE
# REQUIRES: language / add
# ==================================================

@router.post(
    "/languages",
    status_code=status.HTTP_201_CREATED,
)
def create_new_language(

    request: LanguageCreateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("language", "add")
    ),
):

    if language_exists(
        db,
        request.language,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Language already exists",
        )

    language = create_language(
        db,
        request.language,
    )

    return {

        "message": "Language created successfully",

        "language": language_to_dict(language),
    }


# ==================================================
# UPDATE LANGUAGE
# REQUIRES: language / edit
# ==================================================

@router.put(
    "/languages/{language_id}",
)
def update_existing_language(

    language_id: int,

    request: LanguageUpdateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("language", "edit")
    ),
):

    existing = get_language_by_id(
        db,
        language_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Language not found",
        )

    if language_exists_for_other(
        db,
        request.language,
        language_id,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Language already exists",
        )

    language = update_language(
        db,
        language_id,
        request.language,
    )

    return {

        "message": "Language updated successfully",

        "language": language_to_dict(language),
    }


# ==================================================
# DELETE LANGUAGE
# REQUIRES: language / delete
# HARD DELETE (Phase 2)
# ==================================================

@router.delete(
    "/languages/{language_id}",
)
def delete_language(

    language_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("language", "delete")
    ),
):

    existing = get_language_by_id(
        db,
        language_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Language not found",
        )

    hard_delete_language(
        db,
        language_id,
    )

    return {

        "message": "Language deleted successfully",
    }


# ==================================================
# GET REQUIREMENTS
# ==================================================

@router.get(
    "/requirements",
)
def list_requirements(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    requirements = get_requirements(
        db
    )

    return {

        "total": len(requirements),

        "requirements": [

            requirement_to_dict(requirement)

            for requirement in requirements
        ],
    }


# ==================================================
# GET REQUIREMENT BY ID
# ==================================================

@router.get(
    "/requirements/{requirement_id}",
)
def get_requirement(

    requirement_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    requirement = get_requirement_by_id(
        db,
        requirement_id,
    )

    if not requirement:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Requirement not found",
        )

    return {

        "requirement": requirement_to_dict(requirement),
    }


# ==================================================
# CREATE REQUIREMENT
# REQUIRES: requirement / add
# ==================================================

@router.post(
    "/requirements",
    status_code=status.HTTP_201_CREATED,
)
def create_new_requirement(

    request: RequirementCreateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("requirement", "add")
    ),
):

    if requirement_exists(
        db,
        request.requirement,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Requirement already exists",
        )

    requirement = create_requirement(
        db,
        request.requirement,
    )

    return {

        "message": "Requirement created successfully",

        "requirement": requirement_to_dict(requirement),
    }


# ==================================================
# UPDATE REQUIREMENT
# REQUIRES: requirement / edit
# ==================================================

@router.put(
    "/requirements/{requirement_id}",
)
def update_existing_requirement(

    requirement_id: int,

    request: RequirementUpdateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("requirement", "edit")
    ),
):

    existing = get_requirement_by_id(
        db,
        requirement_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Requirement not found",
        )

    if requirement_exists_for_other(
        db,
        request.requirement,
        requirement_id,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Requirement already exists",
        )

    requirement = update_requirement(
        db,
        requirement_id,
        request.requirement,
    )

    return {

        "message": "Requirement updated successfully",

        "requirement": requirement_to_dict(requirement),
    }


# ==================================================
# DELETE REQUIREMENT
# REQUIRES: requirement / delete
# HARD DELETE (Phase 2)
# ==================================================

@router.delete(
    "/requirements/{requirement_id}",
)
def delete_requirement(

    requirement_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("requirement", "delete")
    ),
):

    existing = get_requirement_by_id(
        db,
        requirement_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Requirement not found",
        )

    hard_delete_requirement(
        db,
        requirement_id,
    )

    return {

        "message": "Requirement deleted successfully",
    }



# ==================================================
# GET ADVERTISING MEDIAS
# ==================================================

@router.get(
    "/advertising-medias",
)
def list_advertising_medias(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    advertising_medias = get_advertising_medias(
        db
    )

    return {

        "total": len(advertising_medias),

        "advertising_medias": [

            advertising_media_to_dict(item)

            for item in advertising_medias
        ],
    }


# ==================================================
# GET ADVERTISING MEDIA BY ID
# ==================================================

@router.get(
    "/advertising-medias/{advertising_media_id}",
)
def get_advertising_media(

    advertising_media_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    advertising_media = get_advertising_media_by_id(
        db,
        advertising_media_id,
    )

    if not advertising_media:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Advertising media not found",
        )

    return {

        "advertising_media": advertising_media_to_dict(
            advertising_media
        ),
    }


# ==================================================
# CREATE ADVERTISING MEDIA
# REQUIRES: advertising_media / add
# ==================================================

@router.post(
    "/advertising-medias",
    status_code=status.HTTP_201_CREATED,
)
def create_new_advertising_media(

    request: AdvertisingMediaCreateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("advertising_media", "add")
    ),
):

    if advertising_media_exists(
        db,
        request.advertising_media,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Advertising media already exists",
        )

    advertising_media = create_advertising_media(
        db,
        request.advertising_media,
    )

    return {

        "message": "Advertising media created successfully",

        "advertising_media": advertising_media_to_dict(
            advertising_media
        ),
    }


# ==================================================
# UPDATE ADVERTISING MEDIA
# REQUIRES: advertising_media / edit
# ==================================================

@router.put(
    "/advertising-medias/{advertising_media_id}",
)
def update_existing_advertising_media(

    advertising_media_id: int,

    request: AdvertisingMediaUpdateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("advertising_media", "edit")
    ),
):

    existing = get_advertising_media_by_id(
        db,
        advertising_media_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Advertising media not found",
        )

    if advertising_media_exists_for_other(
        db,
        request.advertising_media,
        advertising_media_id,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Advertising media already exists",
        )

    advertising_media = update_advertising_media(
        db,
        advertising_media_id,
        request.advertising_media,
    )

    return {

        "message": "Advertising media updated successfully",

        "advertising_media": advertising_media_to_dict(
            advertising_media
        ),
    }


# ==================================================
# DELETE ADVERTISING MEDIA
# REQUIRES: advertising_media / delete
# HARD DELETE (Phase 2)
# ==================================================

@router.delete(
    "/advertising-medias/{advertising_media_id}",
)
def delete_advertising_media(

    advertising_media_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("advertising_media", "delete")
    ),
):

    existing = get_advertising_media_by_id(
        db,
        advertising_media_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Advertising media not found",
        )

    hard_delete_advertising_media(
        db,
        advertising_media_id,
    )

    return {

        "message": "Advertising media deleted successfully",
    }



# ==================================================
# GET ADVERTISING PURPOSES
# ==================================================

@router.get(
    "/advertising-purposes",
)
def list_advertising_purposes(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    advertising_purposes = get_advertising_purposes(
        db
    )

    return {

        "total": len(advertising_purposes),

        "advertising_purposes": [

            advertising_purpose_to_dict(item)

            for item in advertising_purposes
        ],
    }


# ==================================================
# GET ADVERTISING PURPOSE BY ID
# ==================================================

@router.get(
    "/advertising-purposes/{advertising_purpose_id}",
)
def get_advertising_purpose(

    advertising_purpose_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    advertising_purpose = get_advertising_purpose_by_id(
        db,
        advertising_purpose_id,
    )

    if not advertising_purpose:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Advertising purpose not found",
        )

    return {

        "advertising_purpose": advertising_purpose_to_dict(
            advertising_purpose
        ),
    }


# ==================================================
# CREATE ADVERTISING PURPOSE
# REQUIRES: advertising_purpose / add
# ==================================================

@router.post(
    "/advertising-purposes",
    status_code=status.HTTP_201_CREATED,
)
def create_new_advertising_purpose(

    request: AdvertisingPurposeCreateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("advertising_purpose", "add")
    ),
):

    if advertising_purpose_exists(
        db,
        request.advertising_purpose,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Advertising purpose already exists",
        )

    advertising_purpose = create_advertising_purpose(
        db,
        request.advertising_purpose,
    )

    return {

        "message": "Advertising purpose created successfully",

        "advertising_purpose": advertising_purpose_to_dict(
            advertising_purpose
        ),
    }


# ==================================================
# UPDATE ADVERTISING PURPOSE
# REQUIRES: advertising_purpose / edit
# ==================================================

@router.put(
    "/advertising-purposes/{advertising_purpose_id}",
)
def update_existing_advertising_purpose(

    advertising_purpose_id: int,

    request: AdvertisingPurposeUpdateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("advertising_purpose", "edit")
    ),
):

    existing = get_advertising_purpose_by_id(
        db,
        advertising_purpose_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Advertising purpose not found",
        )

    if advertising_purpose_exists_for_other(
        db,
        request.advertising_purpose,
        advertising_purpose_id,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Advertising purpose already exists",
        )

    advertising_purpose = update_advertising_purpose(
        db,
        advertising_purpose_id,
        request.advertising_purpose,
    )

    return {

        "message": "Advertising purpose updated successfully",

        "advertising_purpose": advertising_purpose_to_dict(
            advertising_purpose
        ),
    }


# ==================================================
# DELETE ADVERTISING PURPOSE
# REQUIRES: advertising_purpose / delete
# HARD DELETE (Phase 2)
# ==================================================

@router.delete(
    "/advertising-purposes/{advertising_purpose_id}",
)
def delete_advertising_purpose(

    advertising_purpose_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("advertising_purpose", "delete")
    ),
):

    existing = get_advertising_purpose_by_id(
        db,
        advertising_purpose_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Advertising purpose not found",
        )

    hard_delete_advertising_purpose(
        db,
        advertising_purpose_id,
    )

    return {

        "message": "Advertising purpose deleted successfully",
    }



# ==================================================
# GET OFFICE TYPES
# ==================================================

@router.get(
    "/office-types",
)
def list_office_types(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    office_types = get_office_types(
        db
    )

    return {

        "total": len(office_types),

        "office_types": [

            office_type_to_dict(item)

            for item in office_types
        ],
    }


# ==================================================
# GET OFFICE TYPE BY ID
# ==================================================

@router.get(
    "/office-types/{office_type_id}",
)
def get_office_type(

    office_type_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    office_type = get_office_type_by_id(
        db,
        office_type_id,
    )

    if not office_type:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Office type not found",
        )

    return {

        "office_type": office_type_to_dict(
            office_type
        ),
    }


# ==================================================
# CREATE OFFICE TYPE
# REQUIRES: office_type / add
# ==================================================

@router.post(
    "/office-types",
    status_code=status.HTTP_201_CREATED,
)
def create_new_office_type(

    request: OfficeTypeCreateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("office_type", "add")
    ),
):

    if office_type_exists(
        db,
        request.office_type,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Office type already exists",
        )

    office_type = create_office_type(
        db,
        request.office_type,
    )

    return {

        "message": "Office type created successfully",

        "office_type": office_type_to_dict(
            office_type
        ),
    }


# ==================================================
# UPDATE OFFICE TYPE
# REQUIRES: office_type / edit
# ==================================================

@router.put(
    "/office-types/{office_type_id}",
)
def update_existing_office_type(

    office_type_id: int,

    request: OfficeTypeUpdateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("office_type", "edit")
    ),
):

    existing = get_office_type_by_id(
        db,
        office_type_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Office type not found",
        )

    if office_type_exists_for_other(
        db,
        request.office_type,
        office_type_id,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Office type already exists",
        )

    office_type = update_office_type(
        db,
        office_type_id,
        request.office_type,
    )

    return {

        "message": "Office type updated successfully",

        "office_type": office_type_to_dict(
            office_type
        ),
    }


# ==================================================
# DELETE OFFICE TYPE
# REQUIRES: office_type / delete
# HARD DELETE (Phase 2)
# ==================================================

@router.delete(
    "/office-types/{office_type_id}",
)
def delete_office_type(

    office_type_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("office_type", "delete")
    ),
):

    existing = get_office_type_by_id(
        db,
        office_type_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Office type not found",
        )

    hard_delete_office_type(
        db,
        office_type_id,
    )

    return {

        "message": "Office type deleted successfully",
    }



# ==================================================
# GET MEETING LOCATIONS
# ==================================================

@router.get(
    "/meeting-locations",
)
def list_meeting_locations(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    meeting_locations = get_meeting_locations(
        db
    )

    return {

        "total": len(meeting_locations),

        "meeting_locations": [

            meeting_location_to_dict(item)

            for item in meeting_locations
        ],
    }


# ==================================================
# GET MEETING LOCATION BY ID
# ==================================================

@router.get(
    "/meeting-locations/{meeting_location_id}",
)
def get_meeting_location(

    meeting_location_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    meeting_location = get_meeting_location_by_id(
        db,
        meeting_location_id,
    )

    if not meeting_location:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Meeting location not found",
        )

    return {

        "meeting_location": meeting_location_to_dict(
            meeting_location
        ),
    }


# ==================================================
# CREATE MEETING LOCATION
# REQUIRES: meeting_location / add
# ==================================================

@router.post(
    "/meeting-locations",
    status_code=status.HTTP_201_CREATED,
)
def create_new_meeting_location(

    request: MeetingLocationCreateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("meeting_location", "add")
    ),
):

    if meeting_location_exists(
        db,
        request.meeting_location,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Meeting location already exists",
        )

    meeting_location = create_meeting_location(
        db,
        request.meeting_location,
        fk_user_id=resolve_legacy_user_id(db, current_user["pkid"]),
    )

    return {

        "message": "Meeting location created successfully",

        "meeting_location": meeting_location_to_dict(
            meeting_location
        ),
    }


# ==================================================
# UPDATE MEETING LOCATION
# REQUIRES: meeting_location / edit
# ==================================================

@router.put(
    "/meeting-locations/{meeting_location_id}",
)
def update_existing_meeting_location(

    meeting_location_id: int,

    request: MeetingLocationUpdateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("meeting_location", "edit")
    ),
):

    existing = get_meeting_location_by_id(
        db,
        meeting_location_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Meeting location not found",
        )

    if meeting_location_exists_for_other(
        db,
        request.meeting_location,
        meeting_location_id,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Meeting location already exists",
        )

    meeting_location = update_meeting_location(
        db,
        meeting_location_id,
        request.meeting_location,
    )

    return {

        "message": "Meeting location updated successfully",

        "meeting_location": meeting_location_to_dict(
            meeting_location
        ),
    }


# ==================================================
# DELETE MEETING LOCATION
# REQUIRES: meeting_location / delete
# HARD DELETE (Phase 2)
# ==================================================

@router.delete(
    "/meeting-locations/{meeting_location_id}",
)
def delete_meeting_location(

    meeting_location_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("meeting_location", "delete")
    ),
):

    existing = get_meeting_location_by_id(
        db,
        meeting_location_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Meeting location not found",
        )

    hard_delete_meeting_location(
        db,
        meeting_location_id,
    )

    return {

        "message": "Meeting location deleted successfully",
    }



# ==================================================
# GET KSAs
# ==================================================

@router.get(
    "/ksas",
)
def list_ksas(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    ksas = get_ksas(
        db
    )

    return {

        "total": len(ksas),

        "ksas": [

            ksa_to_dict(item)

            for item in ksas
        ],
    }


# ==================================================
# GET KSA BY ID
# ==================================================

@router.get(
    "/ksas/{ksa_id}",
)
def get_ksa(

    ksa_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    ksa = get_ksa_by_id(
        db,
        ksa_id,
    )

    if not ksa:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="KSA not found",
        )

    return {

        "ksa": ksa_to_dict(
            ksa
        ),
    }


# ==================================================
# CREATE KSA
# REQUIRES: ksa / add
# ==================================================

@router.post(
    "/ksas",
    status_code=status.HTTP_201_CREATED,
)
def create_new_ksa(

    request: KSACreateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("ksa", "add")
    ),
):

    if ksa_exists(
        db,
        request.ksa,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="KSA already exists",
        )

    ksa = create_ksa(
        db,
        request.ksa,
    )

    return {

        "message": "KSA created successfully",

        "ksa": ksa_to_dict(
            ksa
        ),
    }


# ==================================================
# UPDATE KSA
# REQUIRES: ksa / edit
# ==================================================

@router.put(
    "/ksas/{ksa_id}",
)
def update_existing_ksa(

    ksa_id: int,

    request: KSAUpdateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("ksa", "edit")
    ),
):

    existing = get_ksa_by_id(
        db,
        ksa_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="KSA not found",
        )

    if ksa_exists_for_other(
        db,
        request.ksa,
        ksa_id,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="KSA already exists",
        )

    ksa = update_ksa(
        db,
        ksa_id,
        request.ksa,
    )

    return {

        "message": "KSA updated successfully",

        "ksa": ksa_to_dict(
            ksa
        ),
    }


# ==================================================
# DELETE KSA
# REQUIRES: ksa / delete
# HARD DELETE (Phase 2)
# ==================================================

@router.delete(
    "/ksas/{ksa_id}",
)
def delete_ksa(

    ksa_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("ksa", "delete")
    ),
):

    existing = get_ksa_by_id(
        db,
        ksa_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="KSA not found",
        )

    hard_delete_ksa(
        db,
        ksa_id,
    )

    return {

        "message": "KSA deleted successfully",
    }



# ==================================================
# GET KSA CATEGORIES
# ==================================================

@router.get(
    "/ksa-categories",
)
def list_ksa_categories(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    ksa_categories = get_ksa_categories(
        db
    )

    return {

        "total": len(ksa_categories),

        "ksa_categories": [

            ksa_category_to_dict(item)

            for item in ksa_categories
        ],
    }


# ==================================================
# GET KSA CATEGORY BY ID
# ==================================================

@router.get(
    "/ksa-categories/{ksa_category_id}",
)
def get_ksa_category(

    ksa_category_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    ksa_category = get_ksa_category_by_id(
        db,
        ksa_category_id,
    )

    if not ksa_category:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="KSA category not found",
        )

    return {

        "ksa_category": ksa_category_to_dict(
            ksa_category
        ),
    }


# ==================================================
# CREATE KSA CATEGORY
# REQUIRES: ksa_category / add
# ==================================================

@router.post(
    "/ksa-categories",
    status_code=status.HTTP_201_CREATED,
)
def create_new_ksa_category(

    request: KSACategoryCreateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("ksa_category", "add")
    ),
):

    if ksa_category_exists(
        db,
        request.ksa_category,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="KSA category already exists",
        )

    ksa_category = create_ksa_category(
        db,
        request.ksa_category,
    )

    return {

        "message": "KSA category created successfully",

        "ksa_category": ksa_category_to_dict(
            ksa_category
        ),
    }


# ==================================================
# UPDATE KSA CATEGORY
# REQUIRES: ksa_category / edit
# ==================================================

@router.put(
    "/ksa-categories/{ksa_category_id}",
)
def update_existing_ksa_category(

    ksa_category_id: int,

    request: KSACategoryUpdateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("ksa_category", "edit")
    ),
):

    existing = get_ksa_category_by_id(
        db,
        ksa_category_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="KSA category not found",
        )

    if ksa_category_exists_for_other(
        db,
        request.ksa_category,
        ksa_category_id,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="KSA category already exists",
        )

    ksa_category = update_ksa_category(
        db,
        ksa_category_id,
        request.ksa_category,
    )

    return {

        "message": "KSA category updated successfully",

        "ksa_category": ksa_category_to_dict(
            ksa_category
        ),
    }


# ==================================================
# DELETE KSA CATEGORY
# REQUIRES: ksa_category / delete
# HARD DELETE (Phase 2)
# ==================================================

@router.delete(
    "/ksa-categories/{ksa_category_id}",
)
def delete_ksa_category(

    ksa_category_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("ksa_category", "delete")
    ),
):

    existing = get_ksa_category_by_id(
        db,
        ksa_category_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="KSA category not found",
        )

    hard_delete_ksa_category(
        db,
        ksa_category_id,
    )

    return {

        "message": "KSA category deleted successfully",
    }



# ==================================================
# GET POSITION GRADES
# ==================================================

@router.get(
    "/position-grades",
)
def list_position_grades(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    position_grades = get_position_grades(
        db
    )

    return {

        "total": len(position_grades),

        "position_grades": [

            position_grade_to_dict(item)

            for item in position_grades
        ],
    }


# ==================================================
# GET POSITION GRADE BY ID
# ==================================================

@router.get(
    "/position-grades/{position_grade_id}",
)
def get_position_grade(

    position_grade_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    position_grade = get_position_grade_by_id(
        db,
        position_grade_id,
    )

    if not position_grade:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Position grade not found",
        )

    return {

        "position_grade": position_grade_to_dict(
            position_grade
        ),
    }


# ==================================================
# CREATE POSITION GRADE
# REQUIRES: position_grade / add
# ==================================================

@router.post(
    "/position-grades",
    status_code=status.HTTP_201_CREATED,
)
def create_new_position_grade(

    request: PositionGradeCreateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("position_grade", "add")
    ),
):

    if request.minimum_pay > request.maximum_pay:

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Minimum pay cannot be greater than maximum pay",
        )

    if position_grade_exists(
        db,
        request.position_grade,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Position grade already exists",
        )

    position_grade = create_position_grade(
        db,
        request.position_grade,
        request.minimum_pay,
        request.maximum_pay,
    )

    return {

        "message": "Position grade created successfully",

        "position_grade": position_grade_to_dict(
            position_grade
        ),
    }


# ==================================================
# UPDATE POSITION GRADE
# REQUIRES: position_grade / edit
# ==================================================

@router.put(
    "/position-grades/{position_grade_id}",
)
def update_existing_position_grade(

    position_grade_id: int,

    request: PositionGradeUpdateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("position_grade", "edit")
    ),
):

    existing = get_position_grade_by_id(
        db,
        position_grade_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Position grade not found",
        )

    if request.minimum_pay > request.maximum_pay:

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Minimum pay cannot be greater than maximum pay",
        )

    if position_grade_exists_for_other(
        db,
        request.position_grade,
        position_grade_id,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Position grade already exists",
        )

    position_grade = update_position_grade(
        db,
        position_grade_id,
        request.position_grade,
        request.minimum_pay,
        request.maximum_pay,
    )

    return {

        "message": "Position grade updated successfully",

        "position_grade": position_grade_to_dict(
            position_grade
        ),
    }


# ==================================================
# DELETE POSITION GRADE
# REQUIRES: position_grade / delete
# HARD DELETE (Phase 2)
# ==================================================

@router.delete(
    "/position-grades/{position_grade_id}",
)
def delete_position_grade(

    position_grade_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("position_grade", "delete")
    ),
):

    existing = get_position_grade_by_id(
        db,
        position_grade_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Position grade not found",
        )

    hard_delete_position_grade(
        db,
        position_grade_id,
    )

    return {

        "message": "Position grade deleted successfully",
    }



# ==================================================
# GET ROLE IN OFFENSES
# ==================================================

@router.get(
    "/role-in-offenses",
)
def list_role_in_offenses(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    items = get_role_in_offenses(
        db
    )

    return {

        "total": len(items),

        "role_in_offenses": [

            role_in_offense_to_dict(item)

            for item in items
        ],
    }


# ==================================================
# GET ROLE IN OFFENSE BY ID
# ==================================================

@router.get(
    "/role-in-offenses/{role_in_offense_id}",
)
def get_role_in_offense(

    role_in_offense_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    item = get_role_in_offense_by_id(
        db,
        role_in_offense_id,
    )

    if not item:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Role in offense not found",
        )

    return {

        "role_in_offense": role_in_offense_to_dict(
            item
        ),
    }


# ==================================================
# CREATE ROLE IN OFFENSE
# REQUIRES: role_in_offense / add
# ==================================================

@router.post(
    "/role-in-offenses",
    status_code=status.HTTP_201_CREATED,
)
def create_new_role_in_offense(

    request: RoleInOffenseCreateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("role_in_offense", "add")
    ),
):

    if request.minimum_penalty > request.maximum_penalty:

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Minimum penalty cannot be greater than maximum penalty",
        )

    if role_in_offense_exists(
        db,
        request.role_in_offense,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Role in offense already exists",
        )

    item = create_role_in_offense(
        db,
        request.role_in_offense,
        request.minimum_penalty,
        request.maximum_penalty,
    )

    return {

        "message": "Role in offense created successfully",

        "role_in_offense": role_in_offense_to_dict(
            item
        ),
    }


# ==================================================
# UPDATE ROLE IN OFFENSE
# REQUIRES: role_in_offense / edit
# ==================================================

@router.put(
    "/role-in-offenses/{role_in_offense_id}",
)
def update_existing_role_in_offense(

    role_in_offense_id: int,

    request: RoleInOffenseUpdateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("role_in_offense", "edit")
    ),
):

    existing = get_role_in_offense_by_id(
        db,
        role_in_offense_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Role in offense not found",
        )

    if request.minimum_penalty > request.maximum_penalty:

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Minimum penalty cannot be greater than maximum penalty",
        )

    if role_in_offense_exists_for_other(
        db,
        request.role_in_offense,
        role_in_offense_id,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Role in offense already exists",
        )

    item = update_role_in_offense(
        db,
        role_in_offense_id,
        request.role_in_offense,
        request.minimum_penalty,
        request.maximum_penalty,
    )

    return {

        "message": "Role in offense updated successfully",

        "role_in_offense": role_in_offense_to_dict(
            item
        ),
    }


# ==================================================
# DELETE ROLE IN OFFENSE
# REQUIRES: role_in_offense / delete
# HARD DELETE (Phase 2)
# ==================================================

@router.delete(
    "/role-in-offenses/{role_in_offense_id}",
)
def delete_role_in_offense(

    role_in_offense_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("role_in_offense", "delete")
    ),
):

    existing = get_role_in_offense_by_id(
        db,
        role_in_offense_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Role in offense not found",
        )

    hard_delete_role_in_offense(
        db,
        role_in_offense_id,
    )

    return {

        "message": "Role in offense deleted successfully",
    }



# ==================================================
# SALARY — NATURE OF WORK
# ==================================================

@router.get("/nature-of-works")
def list_nature_of_works(
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    items = get_nature_of_works(db)
    return {
        "total": len(items),
        "nature_of_works": [nature_of_work_to_dict(i) for i in items],
    }


@router.get("/nature-of-works/{item_id}")
def get_nature_of_work(
    item_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    item = get_nature_of_work_by_id(db, item_id)
    if not item:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Nature of work not found",
        )
    return {"nature_of_work": nature_of_work_to_dict(item)}


@router.post("/nature-of-works", status_code=status.HTTP_201_CREATED)
def create_new_nature_of_work(
    request: NatureOfWorkCreateRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_nature_of_work", "add")
    ),
):
    if nature_of_work_exists(db, request.nature_of_work):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Nature of work already exists",
        )
    item = create_nature_of_work(
        db,
        request.nature_of_work,
        fk_user_id=resolve_legacy_user_id(db, current_user["pkid"]),
    )
    return {
        "message": "Nature of work created successfully",
        "nature_of_work": nature_of_work_to_dict(item),
    }


@router.put("/nature-of-works/{item_id}")
def update_existing_nature_of_work(
    item_id: int,
    request: NatureOfWorkUpdateRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_nature_of_work", "edit")
    ),
):
    existing = get_nature_of_work_by_id(db, item_id)
    if not existing:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Nature of work not found",
        )
    if nature_of_work_exists_for_other(db, request.nature_of_work, item_id):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Nature of work already exists",
        )
    item = update_nature_of_work(db, item_id, request.nature_of_work)
    return {
        "message": "Nature of work updated successfully",
        "nature_of_work": nature_of_work_to_dict(item),
    }


@router.delete("/nature-of-works/{item_id}")
def delete_nature_of_work_route(
    item_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_nature_of_work", "delete")
    ),
):
    existing = get_nature_of_work_by_id(db, item_id)
    if not existing:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Nature of work not found",
        )
    hard_delete_nature_of_work(db, item_id)
    return {"message": "Nature of work deleted successfully"}


# ==================================================
# SALARY — SCHEDULE TYPE
# ==================================================

@router.get("/schedule-types")
def list_schedule_types(
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    items = get_schedule_types(db)
    return {
        "total": len(items),
        "schedule_types": [schedule_type_to_dict(i) for i in items],
    }


@router.get("/schedule-types/{item_id}")
def get_schedule_type(
    item_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    item = get_schedule_type_by_id(db, item_id)
    if not item:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Schedule type not found",
        )
    return {"schedule_type": schedule_type_to_dict(item)}


@router.post("/schedule-types", status_code=status.HTTP_201_CREATED)
def create_new_schedule_type(
    request: ScheduleTypeCreateRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_schedule_type", "add")
    ),
):
    if schedule_type_exists(db, request.schedule_type):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Schedule type already exists",
        )
    item = create_schedule_type(
        db,
        request.schedule_type,
        fk_user_id=resolve_legacy_user_id(db, current_user["pkid"]),
    )
    return {
        "message": "Schedule type created successfully",
        "schedule_type": schedule_type_to_dict(item),
    }


@router.put("/schedule-types/{item_id}")
def update_existing_schedule_type(
    item_id: int,
    request: ScheduleTypeUpdateRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_schedule_type", "edit")
    ),
):
    existing = get_schedule_type_by_id(db, item_id)
    if not existing:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Schedule type not found",
        )
    if schedule_type_exists_for_other(db, request.schedule_type, item_id):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Schedule type already exists",
        )
    item = update_schedule_type(db, item_id, request.schedule_type)
    return {
        "message": "Schedule type updated successfully",
        "schedule_type": schedule_type_to_dict(item),
    }


@router.delete("/schedule-types/{item_id}")
def delete_schedule_type_route(
    item_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_schedule_type", "delete")
    ),
):
    existing = get_schedule_type_by_id(db, item_id)
    if not existing:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Schedule type not found",
        )
    hard_delete_schedule_type(db, item_id)
    return {"message": "Schedule type deleted successfully"}


# ==================================================
# SALARY — RELIGION
# ==================================================

@router.get("/religions")
def list_religions(
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    items = get_religions(db)
    return {
        "total": len(items),
        "religions": [religion_to_dict(i) for i in items],
    }


@router.get("/religions/{item_id}")
def get_religion(
    item_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    item = get_religion_by_id(db, item_id)
    if not item:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Religion not found",
        )
    return {"religion": religion_to_dict(item)}


@router.post("/religions", status_code=status.HTTP_201_CREATED)
def create_new_religion(
    request: ReligionCreateRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_religion", "add")
    ),
):
    if religion_exists(db, request.religion):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Religion already exists",
        )
    item = create_religion(
        db,
        request.religion,
        fk_user_id=resolve_legacy_user_id(db, current_user["pkid"]),
    )
    return {
        "message": "Religion created successfully",
        "religion": religion_to_dict(item),
    }


@router.put("/religions/{item_id}")
def update_existing_religion(
    item_id: int,
    request: ReligionUpdateRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_religion", "edit")
    ),
):
    existing = get_religion_by_id(db, item_id)
    if not existing:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Religion not found",
        )
    if religion_exists_for_other(db, request.religion, item_id):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Religion already exists",
        )
    item = update_religion(db, item_id, request.religion)
    return {
        "message": "Religion updated successfully",
        "religion": religion_to_dict(item),
    }


@router.delete("/religions/{item_id}")
def delete_religion_route(
    item_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_religion", "delete")
    ),
):
    existing = get_religion_by_id(db, item_id)
    if not existing:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Religion not found",
        )
    hard_delete_religion(db, item_id)
    return {"message": "Religion deleted successfully"}


# ==================================================
# SALARY — CASTES
# ==================================================

@router.get("/castes")
def list_castes(
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    items = get_castes(db)
    return {
        "total": len(items),
        "castes": [caste_to_dict(i) for i in items],
    }


@router.get("/castes/{item_id}")
def get_caste(
    item_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    item = get_caste_by_id(db, item_id)
    if not item:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Caste not found",
        )
    return {"caste": caste_to_dict(item)}


@router.post("/castes", status_code=status.HTTP_201_CREATED)
def create_new_caste(
    request: CasteCreateRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_caste", "add")
    ),
):
    if caste_exists(db, request.caste):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Caste already exists",
        )
    item = create_caste(
        db,
        request.caste,
        fk_user_id=resolve_legacy_user_id(db, current_user["pkid"]),
    )
    return {
        "message": "Caste created successfully",
        "caste": caste_to_dict(item),
    }


@router.put("/castes/{item_id}")
def update_existing_caste(
    item_id: int,
    request: CasteUpdateRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_caste", "edit")
    ),
):
    existing = get_caste_by_id(db, item_id)
    if not existing:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Caste not found",
        )
    if caste_exists_for_other(db, request.caste, item_id):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Caste already exists",
        )
    item = update_caste(db, item_id, request.caste)
    return {
        "message": "Caste updated successfully",
        "caste": caste_to_dict(item),
    }


@router.delete("/castes/{item_id}")
def delete_caste_route(
    item_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_caste", "delete")
    ),
):
    existing = get_caste_by_id(db, item_id)
    if not existing:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Caste not found",
        )
    hard_delete_caste(db, item_id)
    return {"message": "Caste deleted successfully"}


# ==================================================
# SALARY — SKIN TONES
# ==================================================

@router.get("/skin-tones")
def list_skin_tones(
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    items = get_skin_tones(db)
    return {
        "total": len(items),
        "skin_tones": [skin_tone_to_dict(i) for i in items],
    }


@router.get("/skin-tones/{item_id}")
def get_skin_tone(
    item_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    item = get_skin_tone_by_id(db, item_id)
    if not item:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Skin tone not found",
        )
    return {"skin_tone": skin_tone_to_dict(item)}


@router.post("/skin-tones", status_code=status.HTTP_201_CREATED)
def create_new_skin_tone(
    request: SkinToneCreateRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_skin_tone", "add")
    ),
):
    if skin_tone_exists(db, request.colour):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Skin tone already exists",
        )
    item = create_skin_tone(
        db,
        request.colour,
        fk_user_id=resolve_legacy_user_id(db, current_user["pkid"]),
    )
    return {
        "message": "Skin tone created successfully",
        "skin_tone": skin_tone_to_dict(item),
    }


@router.put("/skin-tones/{item_id}")
def update_existing_skin_tone(
    item_id: int,
    request: SkinToneUpdateRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_skin_tone", "edit")
    ),
):
    existing = get_skin_tone_by_id(db, item_id)
    if not existing:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Skin tone not found",
        )
    if skin_tone_exists_for_other(db, request.colour, item_id):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Skin tone already exists",
        )
    item = update_skin_tone(db, item_id, request.colour)
    return {
        "message": "Skin tone updated successfully",
        "skin_tone": skin_tone_to_dict(item),
    }


@router.delete("/skin-tones/{item_id}")
def delete_skin_tone_route(
    item_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_skin_tone", "delete")
    ),
):
    existing = get_skin_tone_by_id(db, item_id)
    if not existing:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Skin tone not found",
        )
    hard_delete_skin_tone(db, item_id)
    return {"message": "Skin tone deleted successfully"}


#  ==================================================
# SALARY — TASK STATUS
# ====#==============================================

@router.get("/task-statuses")
def list_task_statuses(
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    items = get_task_statuses(db)
    return {
        "total": len(items),
        "task_statuses": [task_status_to_dict(i) for i in items],
    }


@router.get("/task-statuses/{item_id}")
def get_task_status(
    item_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    item = get_task_status_by_id(db, item_id)
    if not item:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Task status not found",
        )
    return {"task_status": task_status_to_dict(item)}


@router.post("/task-statuses", status_code=status.HTTP_201_CREATED)
def create_new_task_status(
    request: TaskStatusCreateRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_task_status", "add")
    ),
):
    item = create_task_status(
        db,
        request.status,
        request.finish,
        request.cancel,
        fk_user_id=resolve_legacy_user_id(db, current_user["pkid"]),
    )
    return {
        "message": "Task status created successfully",
        "task_status": task_status_to_dict(item),
    }


@router.put("/task-statuses/{item_id}")
def update_existing_task_status(
    item_id: int,
    request: TaskStatusUpdateRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_task_status", "edit")
    ),
):
    existing = get_task_status_by_id(db, item_id)
    if not existing:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Task status not found",
        )

    item = update_task_status(
        db,
        item_id,
        request.status,
        request.finish,
        request.cancel,
    )
    return {
        "message": "Task status updated successfully",
        "task_status": task_status_to_dict(item),
    }


@router.delete("/task-statuses/{item_id}")
def delete_task_status_route(
    item_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_task_status", "delete")
    ),
):
    existing = get_task_status_by_id(db, item_id)
    if not existing:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Task status not found",
        )
    hard_delete_task_status(db, item_id)
    return {"message": "Task status deleted successfully"}

# ==================================================
# SALARY — EMPLOYEE RELATION
# ==================================================

@router.get("/employee-relations")
def list_employee_relations(
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    items = get_employee_relations(db)
    return {
        "total": len(items),
        "employee_relations": [employee_relation_to_dict(i) for i in items],
    }


@router.get("/employee-relations/{item_id}")
def get_employee_relation(
    item_id: str,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    item = get_employee_relation_by_id(db, item_id)
    if not item:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Employee relation not found",
        )
    return {"employee_relation": employee_relation_to_dict(item)}


@router.post("/employee-relations", status_code=status.HTTP_201_CREATED)
def create_new_employee_relation(
    request: EmployeeRelationCreateRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_employee_relation", "add")
    ),
):
    if employee_relation_exists(db, request.relative_name):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Employee relation already exists",
        )
    item = create_employee_relation(
        db,
        request.relative_name,
        fk_user_id=resolve_legacy_user_id(db, current_user["pkid"]),
    )
    return {
        "message": "Employee relation created successfully",
        "employee_relation": employee_relation_to_dict(item),
    }


@router.put("/employee-relations/{item_id}")
def update_existing_employee_relation(
    item_id: str,
    request: EmployeeRelationUpdateRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_employee_relation", "edit")
    ),
):
    existing = get_employee_relation_by_id(db, item_id)
    if not existing:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Employee relation not found",
        )
    if employee_relation_exists_for_other(db, request.relative_name, item_id):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Employee relation already exists",
        )
    item = update_employee_relation(db, item_id, request.relative_name)
    return {
        "message": "Employee relation updated successfully",
        "employee_relation": employee_relation_to_dict(item),
    }


@router.delete("/employee-relations/{item_id}")
def delete_employee_relation_route(
    item_id: str,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_employee_relation", "delete")
    ),
):
    existing = get_employee_relation_by_id(db, item_id)
    if not existing:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Employee relation not found",
        )
    try:
        hard_delete_employee_relation(db, item_id)
    except IntegrityError:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "This relationship is used by employee relatives "
                "and cannot be deleted"
            ),
        )
    return {"message": "Employee relation deleted successfully"}

# ==================================================
# SALARY — SHIFT TIMING
# ==================================================

@router.get("/shift-timings")
def list_shift_timings(
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    items = get_shift_timings(db)
    return {
        "total": len(items),
        "shift_timings": [shift_timing_to_dict(i) for i in items],
    }


@router.get("/shift-timings/{item_id}")
def get_shift_timing(
    item_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    item = get_shift_timing_by_id(db, item_id)

    if not item:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Shift timing not found",
        )

    return {
        "shift_timing": shift_timing_to_dict(item)
    }


@router.post(
    "/shift-timings",
    status_code=status.HTTP_201_CREATED,
)
def create_new_shift_timing(
    request: ShiftTimingCreateRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_shift_timing", "add")
    ),
):
    item = create_shift_timing(
        db,
        request.pkSTId,
        request.shift,
        request.start_work,
        request.end_work,
        request.total_work,
        request.start_break,
        request.end_break,
        request.total_break,
        fk_user_id=resolve_legacy_user_id(db, current_user["pkid"]),
    )

    return {
        "message": "Shift timing created successfully",
        "shift_timing": shift_timing_to_dict(item),
    }


@router.put("/shift-timings/{item_id}")
def update_existing_shift_timing(
    item_id: int,
    request: ShiftTimingUpdateRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_shift_timing", "edit")
    ),
):
    existing = get_shift_timing_by_id(db, item_id)

    if not existing:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Shift timing not found",
        )

    item = update_shift_timing(
        db,
        item_id,
        request.shift,
        request.start_work,
        request.end_work,
        request.total_work,
        request.start_break,
        request.end_break,
        request.total_break,
    )

    return {
        "message": "Shift timing updated successfully",
        "shift_timing": shift_timing_to_dict(item),
    }


@router.delete("/shift-timings/{item_id}")
def delete_shift_timing_route(
    item_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_shift_timing", "delete")
    ),
):
    existing = get_shift_timing_by_id(db, item_id)

    if not existing:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Shift timing not found",
        )

    delete_shift_timing(db, item_id)

    return {
        "message": "Shift timing deleted successfully"
    }
# ==================================================
# GET MY OWN RIGHTS
# ANY LOGGED-IN USER
# Used by the frontend to decide what to show/hide
# for the person currently logged in.
# ==================================================

@router.get(
    "/user-rights/me",
)
def get_my_rights(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    rights_map = get_user_rights_map(
        db,
        current_user["pkid"],
    )

    return {

        "role": current_user["role"],

        "rights": rights_map,
    }


# ==================================================
# GET A USER'S RIGHTS
# ADMIN ONLY
# Used to populate the User Rights matrix when an
# admin selects a user.
# ==================================================

@router.get(
    "/user-rights/{user_id}",
)
def get_rights_for_user(

    user_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_admin
    ),
):

    target_user = get_user_by_id(
        db,
        user_id,
    )

    if not target_user:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="User not found",
        )

    rights_map = get_user_rights_map(
        db,
        user_id,
    )

    return {

        "user_id": user_id,

        "rights": rights_map,
    }


# ==================================================
# SET A USER'S RIGHTS
# ADMIN ONLY
# Saves the whole matrix for one user in a single
# call, replacing whatever was there before for each
# module included in the request.
# ==================================================

@router.put(
    "/user-rights/{user_id}",
)
def set_rights_for_user(

    user_id: int,

    request: SetUserRightsRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_admin
    ),
):

    target_user = get_user_by_id(
        db,
        user_id,
    )

    if not target_user:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="User not found",
        )

    rights_list = [
        item.model_dump()
        for item in request.rights
    ]

    updated_rows = set_user_rights_bulk(
        db,
        user_id,
        rights_list,
    )

    return {

        "message": "User rights updated successfully",

        "rights": [
            user_right_to_dict(row)
            for row in updated_rows
        ],
    }


# ==================================================
# PORTED FROM route_mine.py — SALARY VALIDATION / NORMALIZATION HELPERS
# ==================================================

SALARY_TEXT_LIMITS = {
    "EmpCode": 30,
    "fkTitId": 5,
    "Employee": 50,
    "fkQualId": 5,
    "PAddress": 255,
    "NAddress": 255,
    "fkDepId": 5,
    "fkDegId": 5,
    "fkBnkId": 10,
    "AccountNo": 20,
    "PFNo": 25,
    "ESICNo": 25,
    "PANNo": 25,
    "BloodGrp": 6,
    "WP": 50,
    "Aadhar": 50,
    "CVCopy": 50,
    "LECopy": 50,
    "UserName": 15,
    "Password": 10,
    "Question": 50,
    "Answer": 50,
    "Ext": 10,
    "fkUserId": 5,
    "LastStatus": 10,
    "RTGS": 20,
    "SAddress": 30,
    "fkSetId": 5,
    "Type": 20,
    "Mark": 50,
    "Experience": 5,
    "Police": 50,
    "AddPolice": 255,
    "ContPolice": 25,
    "Personality1": 50,
    "fkP1DesId": 5,
    "P1Address": 255,
    "P1Contact": 25,
    "Personality2": 50,
    "fkP2DesId": 5,
    "P2Address": 255,
    "P2Contact": 25,
    "fkAcctId": 10,
    "Employment": 3,
}


SALARY_REQUIRED_TEXT_DEFAULTS = {
    "EmpCode": "EMP",
    "Employee": "Employee",
    "PAddress": "N/A",
    "NAddress": "N/A",
    "BloodGrp": "N/A",
    "WP": "N/A",
    "CVCopy": "N/A",
    "LECopy": "N/A",
    "UserName": "employee",
    "Password": "password",
    "Question": "N/A",
    "Answer": "N/A",
    "Ext": "N/A",
    "fkUserId": "ADM01",
    "LastStatus": "Active",
    "SAddress": "N/A",
    "Type": "ALL",
    "Mark": "N/A",
    "Police": "N/A",
    "AddPolice": "N/A",
    "ContPolice": "N/A",
    "Personality1": "N/A",
    "P1Address": "N/A",
    "P1Contact": "N/A",
    "Personality2": "N/A",
    "P2Address": "N/A",
    "P2Contact": "N/A",
    "Employment": "FT",
}


# ==================================================
# NOT NULL COLUMNS WITH NO SENSIBLE PLACEHOLDER
#
# These six are NOT NULL on the real SalEmployee table (confirmed
# against INFORMATION_SCHEMA) but are NOT in the dict above, and
# they can't be -- they each have a strict format validator in
# schema.py, so a placeholder like "N/A" would be rejected before
# it ever reached the DB.
#
# The result was that saving an employee without, say, a PAN sent
# NULL for that column and the INSERT died on the NOT NULL
# constraint. That is what the "TEMP DEBUG" blocks in the create
# and update routes below were added to chase.
#
# Empty string is the right value here: it satisfies NOT NULL,
# it's falsy so _check_salary_unique_fields skips it (many
# employees can share "no PAN on file"), and schema.py's
# validators already treat blank as "not supplied". Unlike "N/A"
# it doesn't pollute the column with text that looks like data.
#
# Applied on BOTH create and update -- an edit that clears a PAN
# hits the same constraint a create does.
# ==================================================

SALARY_NOT_NULL_BLANKABLE = {
    "AccountNo",
    "PFNo",
    "ESICNo",
    "PANNo",
    "Aadhar",
    "RTGS",
}


def _salary_default(
    key: str,
    employee_id: int | None,
):

    if key == "EmpCode":

        return f"EMP{employee_id or ''}"

    if key == "UserName":

        return f"emp{employee_id or ''}"

    return SALARY_REQUIRED_TEXT_DEFAULTS.get(
        key,
        "",
    )


# ==================================================
# STATUTORY IDENTIFIER UNIQUENESS
# PFNo (UAN), PANNo, Aadhar, ESICNo, AccountNo, RTGS (IFSC) must each
# be unique across employees whenever a value is actually provided
# (format is already enforced in schema.py's field validators). This
# check runs on both create AND update -- an edit that reuses another
# employee's PAN, for example, must be rejected exactly the same way
# a create would be.
# ==================================================

SALARY_UNIQUE_FIELD_LABELS = {
    "PFNo": "UAN (PF No.)",
    "PANNo": "PAN",
    "Aadhar": "Aadhar number",
    "ESICNo": "ESIC number",
    "AccountNo": "Bank account number",
    "RTGS": "IFSC code",
    "UserName": "Username",
}


def _check_salary_unique_fields(
    db: Session,
    payload: dict,
    exclude_emp_id: int | None = None,
):

    for field_name, label in SALARY_UNIQUE_FIELD_LABELS.items():

        value = payload.get(field_name)

        if not value:

            continue

        if salary_employee_field_exists(
            db,
            field_name,
            value,
            exclude_emp_id=exclude_emp_id,
        ):

            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=(
                    f"{label} '{value}' is already used by "
                    "another employee."
                ),
            )


# ==================================================
# EMPLOYEE CODE — RE-HIRE PERIOD OVERLAP
# ==================================================
# Ports frmEmployee.ValidateFields' EmpCode/DOJ/DOL block. An EmpCode
# is deliberately reused across a person's separate stints, so this is
# NOT the same shape as SALARY_UNIQUE_FIELD_LABELS above -- two rows
# CAN share an EmpCode, they just can't have overlapping [DOJ, DOL]
# periods, and only one of them may still be open (DOL NULL) at a
# time. doj/dol/employee_name are the EFFECTIVE values for this save
# (existing row's values merged with whatever the request changed),
# not necessarily what's in the raw request payload.
# ==================================================

def _check_employee_code_period(
    db: Session,
    emp_code: str | None,
    doj,
    dol,
    employee_name: str | None,
    exclude_emp_id: int | None = None,
):

    if not emp_code:
        return

    if doj is not None:

        conflict = find_employee_period_overlap(
            db, emp_code, doj, exclude_emp_id=exclude_emp_id
        )

        if conflict:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=(
                    f"Joining Date {doj.strftime('%d/%b/%Y')} is "
                    "overlapping with some existing record."
                ),
            )

    if dol is not None:

        conflict = find_employee_period_overlap(
            db, emp_code, dol, exclude_emp_id=exclude_emp_id
        )

        if conflict:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=(
                    f"Leaving Date {dol.strftime('%d/%b/%Y')} is "
                    "overlapping with some existing record."
                ),
            )

    active = find_employee_active_record(
        db, emp_code, exclude_emp_id=exclude_emp_id
    )

    if active:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                f"Please enter Leaving Date of {employee_name or 'this employee'} "
                f"where Joining Date is "
                f"{active.DOJ.strftime('%d/%b/%Y') if active.DOJ else 'unknown'}."
            ),
        )


# ==================================================
# JOINING / LEAVING DATE vs SALARY STRUCTURE
# ==================================================
# Ports the SalStructure half of dtDOJ_Validating / dtDOL_Validating.
# Edit-mode only, matching the legacy "If Fv.EditMode = True" gate --
# a brand-new employee can't have a salary structure yet, so this is a
# no-op on create regardless.
#
# The legacy handlers also check SalAttendance (MIN/MAX AtDate joined
# through SalStructure) the same way. That half is NOT ported here --
# there's no ORM model for SalAttendance in this codebase yet, and
# guessing its columns from the VB SQL alone felt like the wrong
# tradeoff for a live attendance table. Flagging this so it isn't
# mistaken for full parity: a Leaving/Joining Date that conflicts with
# actual attendance records, rather than a salary structure, will
# currently be accepted here when the legacy form would have rejected
# it.
# ==================================================

def _check_employee_dates_against_structures(
    db: Session,
    emp_id: int,
    doj_changed: bool,
    doj,
    dol_changed: bool,
    dol,
):

    if not (doj_changed or dol_changed):
        return

    min_start, max_revise = get_salary_structure_date_bounds(db, emp_id)

    if dol_changed and dol is not None and max_revise and max_revise > dol:

        formatted = max_revise.strftime("%d/%b/%Y")

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                f"The last salary structure end on {formatted}, then "
                f"enter\\select Leaving Date greater than or equal to {formatted}."
            ),
        )

    if doj_changed and doj is not None and min_start and min_start < doj:

        formatted = min_start.strftime("%d/%b/%Y")

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                f"The first salary structure start from {formatted}, then "
                f"enter\\select Joining Date less than or equal to {formatted}."
            ),
        )


# ==================================================
# PHOTO NORMALIZATION (WRITE SIDE)
#
# The form sends the image as a base64 string. Pydantic's
# `bytes` type coerces that string straight to utf-8 bytes,
# so without this the column ends up holding base64 TEXT
# rather than an image. Decode it here so the DB stores
# real image bytes.
#
# Convention used by the update route:
#   Photo omitted / None -> leave the existing photo alone
#   Photo == "" (empty)  -> explicitly clear the photo
# ==================================================

def _normalize_photo_value(value):

    if value is None:

        return None


    if isinstance(value, memoryview):

        value = value.tobytes()


    if isinstance(value, bytes):

        try:

            value = value.decode("ascii")

        except UnicodeDecodeError:

            # Already real binary — store as-is.
            return value


    if not isinstance(value, str):

        return None


    text = value.strip()

    if not text:

        return None


    if text.startswith("data:"):

        _, _, text = text.partition(",")


    try:

        return base64.b64decode(
            text,
            validate=False,
        )

    except (
        binascii.Error,
        ValueError,
    ):

        return None


def normalize_salary_employee_payload(
    payload: dict,
    include_defaults: bool,
):

    employee_id = payload.get("pkEmpId")

    if "Photo" in payload:

        payload["Photo"] = _normalize_photo_value(
            payload["Photo"]
        )

    if include_defaults:

        for key in SALARY_REQUIRED_TEXT_DEFAULTS:

            if payload.get(key) in (None, ""):

                payload[key] = _salary_default(
                    key,
                    employee_id,
                )

    # Runs on create AND update (outside the include_defaults
    # block) -- see SALARY_NOT_NULL_BLANKABLE above. Only keys
    # actually present are touched, so a partial update never
    # blanks a column the caller didn't mention.
    for key in SALARY_NOT_NULL_BLANKABLE:

        if key in payload and payload[key] is None:

            payload[key] = ""

    for key, max_length in SALARY_TEXT_LIMITS.items():

        value = payload.get(key)

        if value is None:

            continue

        payload[key] = str(value).strip()[:max_length]

    return payload


# ==================================================
# PORTED FROM route_mine.py — ACCOUNT QUICK-CREATE, GENERIC LOOKUPS, USER-ACCOUNT MAP
# ==================================================

# ==================================================
# ACCOUNT QUICK-CREATE
#
# AcctAccount is a real legacy chart-of-accounts table with many
# NOT NULL columns, so the generic one-field lookup creator cannot
# safely insert it. These two endpoints expose only the information
# needed by the Salary Structure picker and fill the remaining legacy
# columns with type-safe defaults.
# ==================================================

@router.get("/accounts/groups")
def list_account_groups(
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    rows = (
        db.query(AcctAccount.fkGrpId)
        .filter(AcctAccount.fkGrpId.isnot(None))
        .distinct()
        .order_by(AcctAccount.fkGrpId.asc())
        .all()
    )

    return {
        "groups": [
            {"value": int(row[0]), "label": f"Account Group {int(row[0])}"}
            for row in rows
            if row[0] is not None
        ]
    }


def _next_account_code(db: Session) -> str:
    rows = db.query(AcctAccount.pkAcctId).all()
    highest = 0
    for row in rows:
        value = str(row[0] or "")
        match = re.fullmatch(r"AC(\d+)", value.upper())
        if match:
            highest = max(highest, int(match.group(1)))
    return f"AC{highest + 1:08d}"


def _account_default_for_column(column, supplied: dict, foreign_keys: dict):
    name = column["name"]

    if name in supplied:
        return supplied[name]

    if name == "pkAcctId":
        return _next_account_code(supplied["_db"])

    if name == "Account":
        return supplied["account"]

    if name == "AcctCode":
        return supplied["acct_code"]

    if name == "fkGrpId":
        return supplied["fk_grp_id"]

    if name == "fkUserId":
        return supplied["legacy_user_id"]

    if name == "Active":
        return True

    if name == "SysDefined":
        return False

    if name == "Sync":
        return "N"

    if name == "LastStatus":
        return "Added"

    type_text = str(column["type"]).lower()

    if "date" in type_text or "time" in type_text:
        return datetime.now()

    if "bit" in type_text:
        return False

    if any(token in type_text for token in ("int", "numeric", "decimal", "money", "float", "real")):
        return 0

    if "binary" in type_text or "image" in type_text:
        return b""

    if column.get("nullable"):
        return None

    # All remaining non-null legacy text columns get an empty string.
    return ""


@router.post("/accounts/quick-create", status_code=status.HTTP_201_CREATED)
def create_account_quick(
    payload: dict = Body(...),
    db: Session = Depends(get_db),
    current_user: dict = Depends(require_permission("lookup_masters", "add")),
):
    account = str(payload.get("account") or "").strip()
    acct_code = str(payload.get("acct_code") or "").strip()
    fk_grp_id = payload.get("fk_grp_id")

    if not account:
        raise HTTPException(status_code=400, detail="Account Name is required.")
    if not acct_code:
        raise HTTPException(status_code=400, detail="Account Code is required.")
    if len(acct_code) > 20:
        raise HTTPException(status_code=400, detail="Account Code cannot exceed 20 characters.")
    if fk_grp_id in (None, ""):
        raise HTTPException(status_code=400, detail="Account Group is required.")

    try:
        fk_grp_id = int(fk_grp_id)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="Account Group must be a valid group ID.")

    duplicate_name = (
        db.query(AcctAccount)
        .filter(AcctAccount.Account.ilike(account))
        .first()
    )
    if duplicate_name:
        raise HTTPException(status_code=400, detail="An account with this name already exists.")

    duplicate_code = (
        db.query(AcctAccount)
        .filter(AcctAccount.AcctCode.ilike(acct_code))
        .first()
    )
    if duplicate_code:
        raise HTTPException(status_code=400, detail="An account with this code already exists.")

    inspector = inspect(db.bind)
    columns = inspector.get_columns("AcctAccount", schema="dbo")
    foreign_keys = inspector.get_foreign_keys("AcctAccount", schema="dbo")

    # Validate the selected group against an existing AcctAccount group
    # before touching the legacy table. This avoids inventing a group ID.
    group_exists = (
        db.query(AcctAccount.fkGrpId)
        .filter(AcctAccount.fkGrpId == fk_grp_id)
        .first()
    )
    if not group_exists:
        raise HTTPException(status_code=400, detail="Selected Account Group does not exist in the current account master.")

    legacy_user_id = resolve_legacy_user_id(db, current_user["pkid"])
    supplied = {
        "account": account,
        "acct_code": acct_code,
        "fk_grp_id": fk_grp_id,
        "legacy_user_id": legacy_user_id,
        "_db": db,
    }

    # Known FK targets. If another non-nullable FK exists in the real
    # table, use its first existing key rather than writing a made-up FK.
    for fk in foreign_keys:
        for local_column in fk.get("constrained_columns") or []:
            if local_column in {"fkGrpId"}:
                continue
            if local_column in supplied:
                continue
            if not local_column.startswith("fk"):
                continue
            referred_table = (fk.get("referred_table") or "").strip()
            referred_schema = fk.get("referred_schema") or "dbo"
            referred_columns = fk.get("referred_columns") or []
            if not referred_table or not referred_columns:
                continue
            referred_column = referred_columns[0]
            row = db.execute(
                text(
                    f"SELECT TOP 1 [{referred_column}] FROM "
                    f"[{referred_schema}].[{referred_table}]"
                )
            ).first()
            if row:
                supplied[local_column] = row[0]

    values = {}
    for column in columns:
        if column.get("computed") or column.get("identity"):
            continue
        name = column["name"]
        if name == "pkAcctId":
            values[name] = _next_account_code(db)
        else:
            values[name] = _account_default_for_column(column, supplied, foreign_keys)

    column_names = list(values.keys())
    sql_columns = ", ".join(f"[{name}]" for name in column_names)
    sql_params = ", ".join(f":p{i}" for i in range(len(column_names)))
    params = {f"p{i}": values[name] for i, name in enumerate(column_names)}

    try:
        db.execute(
            text(
                f"INSERT INTO [dbo].[AcctAccount] ({sql_columns}) "
                f"VALUES ({sql_params})"
            ),
            params,
        )
        db.commit()
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=400, detail=f"Could not create account: {exc}")

    option = {
        "value": values["pkAcctId"],
        "label": account,
    }

    return {
        "message": "Account created.",
        "option": option,
    }


# ==================================================
# GENERIC LOOKUP MASTERS (fk* dropdown + "add new" flow)
# Backs every fk*Id combobox on SalaryEmployeeMaster.jsx (and,
# per the "so many tables to wire later" plan, future master
# forms too) via one config-driven endpoint set instead of 12
# near-identical CRUD blocks. See LOOKUP_TABLE_CONFIG in model.py
# for the table -> slug mapping.
#
# REQUIRES: lookup_masters / view, add (reused across all tables
# rather than a per-table permission, since these are all small
# reference-data masters editable inline from any form).
# ==================================================

def _lookup_config_or_404(table: str) -> dict:

    config = LOOKUP_TABLE_CONFIG.get(table)

    if not config:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Unknown lookup table: {table}",
        )

    return config


@router.get(
    "/lookups/{table}",
)
def list_lookup_rows(

    table: str,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    config = _lookup_config_or_404(table)

    rows = get_lookup_rows(db, config)

    return {

        "table": table,

        "allow_create": config["allow_create"],

        "total": len(rows),

        "options": rows,
    }


@router.post(
    "/lookups/{table}",
    status_code=status.HTTP_201_CREATED,
)
def create_new_lookup_row(

    table: str,

    request: LookupCreateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("lookup_masters", "add")
    ),
):

    config = _lookup_config_or_404(table)

    if not config["allow_create"]:

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                f"New values can't be added to '{table}' "
                "from here."
            ),
        )

    if lookup_label_exists(db, config, request.label):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="This value already exists.",
        )

    try:

        row = create_lookup_row(
            db,
            config,
            request.label,
            resolve_legacy_user_id(db, current_user["pkid"]),
        )

    except ValueError as exc:

        # e.g. resolve_docmas_setid finding no IDSettings row to
        # reference -- a real setup gap, not a schema mismatch.
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(exc),
        ) from exc

    except (DataError, IntegrityError) as exc:

        # Mirrors the SalEmployee debug block — these 12 tables'
        # full schemas are unverified (see model.py comment), so
        # surface the real SQL error on the console immediately
        # rather than guessing.
        print(f"=== LOOKUP INSERT FAILED ({table}) ===")
        print(repr(exc.orig) if hasattr(exc, "orig") else repr(exc))
        traceback.print_exc()
        print("=======================================")

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                f"Could not add this value to '{table}' — it may "
                "not match that table's real schema."
            ),
        ) from exc

    return {

        "message": "Value added successfully",

        "option": lookup_row_to_dict(config, row),
    }


# ==================================================
# USER ACCOUNT MAP
# Admin-only. Maps a JWT app_users.pkid to its legacy
# AppUser.pkUserId code so fkUserId audit columns can reflect
# who's actually logged in (see resolve_legacy_user_id in
# model.py). Until a user has a mapping row, legacy writes fall
# back to LEGACY_FALLBACK_USER_ID.
# ==================================================

@router.get(
    "/user-account-map",
)
def list_account_maps(

    db: Session = Depends(get_db),

    current_user: dict = Depends(require_admin),
):

    mappings = list_user_account_maps(db)

    return {

        "total": len(mappings),

        "mappings": [
            {
                "fkAppUserId": mapping.fkAppUserId,
                "LegacyUserId": mapping.LegacyUserId,
            }
            for mapping in mappings
        ],
    }


@router.put(
    "/user-account-map",
)
def upsert_account_map(

    request: UserAccountMapRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(require_admin),
):

    mapping = upsert_user_account_map(
        db,
        request.fkAppUserId,
        request.LegacyUserId,
    )

    return {

        "message": "Mapping saved successfully",

        "mapping": {
            "fkAppUserId": mapping.fkAppUserId,
            "LegacyUserId": mapping.LegacyUserId,
        },
    }


# ==================================================
# PORTED FROM route_mine.py — SALARY EMPLOYEES, DOCUMENTS/CONTACTS/RELATIONS GRIDS, SALARY STRUCTURES
# ==================================================

# ==================================================
# GET SALARY EMPLOYEES
# NOTE: plain get_current_user for now, same as every
# other GET route — see the TODO in security.py about
# whether list/get should require can_view.
# ==================================================

@router.get(
    "/salary-employees",
)
def list_salary_employees(

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    employees = get_salary_employees(
        db
    )

    return {

        "total": len(employees),

        "employees": [

            salary_employee_to_dict(employee)

            for employee in employees
        ],
    }


# ==================================================
# SEARCH SALARY EMPLOYEES (employee picker)
# Backs fkREmpId / fkW1EmpId / fkW2EmpId — these reference real
# employee rows, not a lookup master, so it's search-only, no
# "add new" affordance. Registered ahead of the {emp_id} route
# below so "search" isn't swallowed as an int path param.
# ==================================================

@router.get(
    "/salary-employees/search",
)
def search_employees(

    q: str = "",

    exclude: int | None = None,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    if not q or len(q.strip()) < 2:

        return {"total": 0, "employees": []}

    employees = search_salary_employees(
        db,
        q.strip(),
        exclude_emp_id=exclude,
    )

    return {

        "total": len(employees),

        "employees": [
            {
                "pkEmpId": employee.pkEmpId,
                "EmpCode": employee.EmpCode,
                "Employee": employee.Employee,
            }
            for employee in employees
        ],
    }


# ==================================================
# GET SALARY EMPLOYEE BY ID
# ==================================================

@router.get(
    "/salary-employees/{emp_id}",
)
def get_salary_employee(

    emp_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    employee = get_salary_employee_by_id(
        db,
        emp_id,
    )

    if not employee:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Employee not found",
        )

    return {

        "employee": salary_employee_to_dict(employee),
    }


# ==================================================
# CREATE SALARY EMPLOYEE
# REQUIRES: salary_employee / add
#
# pkEmpId is NOT auto-generated (the underlying table
# is a plain integer primary key, not a SERIAL column —
# see schema.py), so the caller must supply it. The
# frontend suggests the next free id, but the id is
# still editable, matching the legacy Employee Code
# behaviour shown in the reference screenshots.
# ==================================================

@router.post(
    "/salary-employees",
    status_code=status.HTTP_201_CREATED,
)
def create_new_salary_employee(

    request: SalEmployeeCreateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("salary_employee", "add")
    ),
):

    if get_salary_employee_by_id(
        db,
        request.pkEmpId,
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="An employee with this Employee ID already exists",
        )

    payload = normalize_salary_employee_payload(
        request.model_dump(),
        include_defaults=True,
    )

    # fkUserId is a legacy audit column, not a value the frontend
    # should be trusted to set — always resolve it from the actual
    # logged-in JWT user (see resolve_legacy_user_id in model.py).
    payload["fkUserId"] = resolve_legacy_user_id(
        db,
        current_user["pkid"],
    )

    # fkSetId is auto-assigned, not user-facing -- same category as
    # fkUserId above. The frontend still sends a placeholder value for
    # it (currently the literal "SET01", which doesn't exist as a real
    # IDSettings row), so this MUST be an unconditional overwrite, not
    # a fallback-if-missing check -- the frontend always sends
    # something truthy, so "if not payload.get(...)" never fires and
    # the placeholder sails straight through. Always resolve from a
    # real IDSettings row instead of trusting anything the client sent.
    try:

        payload["fkSetId"] = resolve_default_set_id(db)

    except ValueError as exc:

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(exc),
        ) from exc

    # ValidateFields: Joining Date is a hard requirement on every
    # save, not just something normalize_salary_employee_payload can
    # quietly fill in.
    if payload.get("DOJ") is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Please enter\\select Joining Date.",
        )

    _check_salary_unique_fields(db, payload)

    _check_employee_code_period(
        db,
        payload.get("EmpCode"),
        payload.get("DOJ"),
        payload.get("DOL"),
        payload.get("Employee"),
    )

    try:

        employee = create_salary_employee(
            db,
            payload,
        )

    except (DataError, IntegrityError) as exc:

        # TEMP DEBUG — remove once the real cause is found.
        print("=== SalEmployee INSERT FAILED ===")
        print(repr(exc.orig) if hasattr(exc, "orig") else repr(exc))
        traceback.print_exc()
        print("==================================")

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "Employee could not be saved because one or more "
                "values do not match the SalEmployee table schema."
            ),
        ) from exc

    return {

        "message": "Employee created successfully",

        "employee": salary_employee_to_dict(employee),
    }


# ==================================================
# UPDATE SALARY EMPLOYEE
# REQUIRES: salary_employee / edit
# ==================================================

@router.put(
    "/salary-employees/{emp_id}",
)
def update_existing_salary_employee(

    emp_id: int,

    request: SalEmployeeUpdateRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("salary_employee", "edit")
    ),
):

    existing = get_salary_employee_by_id(
        db,
        emp_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Employee not found",
        )

    update_payload = request.model_dump(
        exclude_unset=True
    )

    # A null Photo means "not supplied" — drop it so an edit
    # can never silently wipe an existing photo. Clearing is
    # done by sending an empty string, which normalization
    # turns into None below.

    if update_payload.get("Photo") is None:

        update_payload.pop(
            "Photo",
            None,
        )

    normalized_payload = normalize_salary_employee_payload(
        update_payload,
        include_defaults=False,
    )

    # Same as create — fkUserId always reflects the actual logged-in
    # JWT user, never whatever (if anything) the frontend sent.
    normalized_payload["fkUserId"] = resolve_legacy_user_id(
        db,
        current_user["pkid"],
    )

    # fkSetId isn't user-facing, but the frontend's edit form builds
    # its payload the same way the create form does, so it can still
    # arrive here as the placeholder "SET01". Same reasoning as
    # create_new_salary_employee above: always overwrite, never a
    # fallback-if-missing check.
    try:

        normalized_payload["fkSetId"] = resolve_default_set_id(db)

    except ValueError as exc:

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(exc),
        ) from exc

    _check_salary_unique_fields(
        db,
        normalized_payload,
        exclude_emp_id=emp_id,
    )

    # ValidateFields runs identically for add and edit -- Joining Date,
    # the EmpCode period-overlap rules, and the SalStructure date
    # bounds are all still enforced when editing. "Effective" values
    # merge whatever this PUT actually changed onto the row as it
    # already stands, since exclude_unset=True means an untouched
    # field simply won't be a key in normalized_payload at all.
    effective_doj = normalized_payload.get("DOJ", existing.DOJ)
    effective_dol = normalized_payload.get("DOL", existing.DOL)
    effective_emp_code = normalized_payload.get("EmpCode", existing.EmpCode)
    effective_name = normalized_payload.get("Employee", existing.Employee)

    if effective_doj is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Please enter\\select Joining Date.",
        )

    _check_employee_code_period(
        db,
        effective_emp_code,
        effective_doj,
        effective_dol,
        effective_name,
        exclude_emp_id=emp_id,
    )

    _check_employee_dates_against_structures(
        db,
        emp_id,
        doj_changed="DOJ" in normalized_payload,
        doj=effective_doj,
        dol_changed="DOL" in normalized_payload,
        dol=effective_dol,
    )

    try:

        employee = update_salary_employee(
            db,
            existing,
            **normalized_payload,
        )

    except (DataError, IntegrityError) as exc:

        # Mirrored from create_new_salary_employee — same reasoning:
        # this block was swallowing the real SQL Server error.
        print("=== SalEmployee UPDATE FAILED ===")
        print(repr(exc.orig) if hasattr(exc, "orig") else repr(exc))
        traceback.print_exc()
        print("==================================")

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "Employee could not be updated because one or more "
                "values do not match the SalEmployee table schema."
            ),
        ) from exc

    return {

        "message": "Employee updated successfully",

        "employee": salary_employee_to_dict(employee),
    }


# ==================================================
# DELETE SALARY EMPLOYEE
# REQUIRES: salary_employee / delete
# ==================================================

@router.delete(
    "/salary-employees/{emp_id}",
)
def delete_existing_salary_employee(

    emp_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("salary_employee", "delete")
    ),
):

    existing = get_salary_employee_by_id(
        db,
        emp_id,
    )

    if not existing:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Employee not found",
        )

    # Child rows (contacts / relatives / documents) FK back
    # to this employee, so they have to go first or the
    # delete is blocked. Document files are removed from
    # disk too, so nothing is orphaned.

    removed_documents = delete_employee_children(
        db,
        emp_id,
    )

    for document in removed_documents:

        stored_name = document.get("DocFile") or ""

        if stored_name:

            (
                _employee_document_dir(emp_id) / stored_name
            ).unlink(
                missing_ok=True
            )

    delete_salary_employee(
        db,
        existing,
    )

    return {

        "message": "Employee deleted successfully",
    }


# ==================================================
# EMPLOYEE DOCUMENT STORAGE
#
# The legacy SalEmpDocuments.DocFile column is only
# nvarchar(100) — it holds a FILE NAME, not the file. So
# uploads are written to disk here and the generated name
# is what goes in the column.
#
# Stored name format: "<12 hex chars>_<original name>",
# truncated to 100 chars with the extension preserved.
# Keeping the original name in the string means the grid
# can show something meaningful without needing an extra
# column the legacy table doesn't have.
# ==================================================

EMPLOYEE_DOCUMENT_ROOT = (
    Path(__file__).resolve().parent
    / "uploads"
    / "employee-documents"
)

DOCUMENT_FILENAME_MAX = 100

MAX_DOCUMENT_BYTES = 10 * 1024 * 1024

ALLOWED_DOCUMENT_EXTENSIONS = {
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".bmp",
    ".webp",
    ".doc",
    ".docx",
    ".xls",
    ".xlsx",
    ".txt",
}


def _employee_document_dir(
    emp_id: int,
) -> Path:

    return EMPLOYEE_DOCUMENT_ROOT / str(emp_id)


def _safe_document_name(
    original_name: str,
) -> str:

    name = Path(
        original_name or "document"
    ).name

    # Strip anything that isn't a safe filename character.
    name = re.sub(
        r"[^A-Za-z0-9._-]",
        "_",
        name,
    ).strip("._") or "document"

    suffix = Path(name).suffix.lower()

    stem = Path(name).stem

    prefix = uuid.uuid4().hex[:12] + "_"

    # Budget: prefix + stem + suffix must fit the column.
    stem_budget = (
        DOCUMENT_FILENAME_MAX
        - len(prefix)
        - len(suffix)
    )

    if stem_budget < 1:

        # Pathological extension — drop the stem entirely.
        return (prefix + suffix)[:DOCUMENT_FILENAME_MAX]

    return prefix + stem[:stem_budget] + suffix


def document_display_name(
    stored_name: str,
) -> str:

    # Reverse of the prefix added above, for display only.

    if not stored_name:

        return ""

    parts = stored_name.split(
        "_",
        1,
    )

    if (
        len(parts) == 2
        and len(parts[0]) == 12
        and re.fullmatch(r"[0-9a-f]{12}", parts[0])
    ):

        return parts[1]

    return stored_name


def employee_document_payload(
    emp_id: int,
    document,
) -> dict:

    data = employee_document_to_dict(document)

    stored_name = data.get("DocFile") or ""

    file_path = (
        _employee_document_dir(emp_id) / stored_name
        if stored_name
        else None
    )

    # `file_exists` lets the UI flag a row whose file went
    # missing on disk, instead of showing a link that
    # silently 404s.

    data["DisplayName"] = document_display_name(
        stored_name
    )

    data["FileExists"] = bool(
        file_path and file_path.is_file()
    )

    data["FileUrl"] = (
        f"/api/salary-employees/{emp_id}"
        f"/documents/{data.get('pkDEmpId')}/file"
    )

    return data


# ==================================================
# CHILD GRID NORMALIZATION
#
# Same idea as normalize_salary_employee_payload: the
# legacy DDL has NOT NULL + tight lengths on these
# columns, so blank grid cells get a default and long
# text gets truncated rather than erroring.
# ==================================================

CONTACT_TEXT_LIMITS = {
    "fkMOCId": 5,
    "Contact": 50,
    "Ext": 10,
}

CONTACT_REQUIRED_DEFAULTS = {
    "fkMOCId": "MOC01",
    "Contact": "N/A",
    "Ext": "N/A",
}

RELATION_TEXT_LIMITS = {
    "RelativeName": 50,
    "fkRelId": 5,
    "fkQuaId": 5,
    "fkSchId": 10,
    "MS": 15,
    "fkDesId": 5,
}

RELATION_REQUIRED_DEFAULTS = {
    "RelativeName": "N/A",
    "fkRelId": "REL01",
    "MS": "N/A",
}

# Columns that are genuinely nullable in the legacy DDL —
# blank stays NULL rather than becoming "".
RELATION_NULLABLE_TEXT = {
    "fkQuaId",
    "fkSchId",
    "fkDesId",
}


def normalize_contact_rows(
    rows: list,
) -> list:

    normalized = []

    for index, row in enumerate(rows):

        data = row.model_dump()

        for key, default in CONTACT_REQUIRED_DEFAULTS.items():

            if data.get(key) in (None, ""):

                data[key] = default

        for key, max_length in CONTACT_TEXT_LIMITS.items():

            value = data.get(key)

            if value is None:

                continue

            data[key] = str(value).strip()[:max_length]

        if data.get("SrNo") in (None, ""):

            data["SrNo"] = index + 1

        normalized.append(data)

    return normalized


def normalize_relation_rows(
    rows: list,
) -> list:

    normalized = []

    for row in rows:

        data = row.model_dump()

        for key, default in RELATION_REQUIRED_DEFAULTS.items():

            if data.get(key) in (None, ""):

                data[key] = default

        for key in RELATION_NULLABLE_TEXT:

            if data.get(key) == "":

                data[key] = None

        for key, max_length in RELATION_TEXT_LIMITS.items():

            value = data.get(key)

            if value is None:

                continue

            data[key] = str(value).strip()[:max_length]

        normalized.append(data)

    return normalized


# ==================================================
# SALARY STRUCTURE
#
# The API field names below intentionally match the existing
# SalaryStructure React screens.  The database model uses the exact
# legacy SalStructureTest column names.  Keeping this translation here
# means the existing ERP routes above remain untouched while Salary
# Structure can be wired end-to-end.
# ==================================================

SAL_STRUCTURE_TYPE_FIELDS = {
    "AllowanceType": "TAllowance",
    "TravelAllowanceType": "TTravelling",
    "HousingAllowanceType": "THousing",
    "DearnessAllowanceType": "TDaily",
    "IncentiveType": "TIncentive",
    "EducationAllowanceType": "TEducation",
    "MedicalAllowanceType": "TMedical",
    "OtherAllowanceType": "TOther",
}

SAL_STRUCTURE_FIELD_MAP = {
    "fkEmpId": "fkEmpId",
    "SalaryStart": "SalStart",
    "Basic": "Basic",
    "BasicType": "BType",
    "DailySalary": "SalDaily",
    "Allowance": "Allowance",
    "TravelAllowance": "Travelling",
    "HousingAllowance": "Housing",
    "DearnessAllowance": "Daily",
    "Incentive": "Incentive",
    "EducationAllowance": "Education",
    "MedicalAllowance": "Medical",
    "OtherAllowance": "Other",
    "OvertimeI": "OTI",
    "OvertimeII": "OTII",
    "OvertimeIType": "TOTI",
    "OvertimeIIType": "TOTII",
    "RestDay1": "RDayI",
    "RestDay2": "RDayII",
    "PaidHoliday": "PH",
    "SickLeave": "SL",
    "PaidCasualLeave": "CL",
    "UnpaidCasualLeave": "UCL",
    "WorkingHoursPerDay": "WH",
    "ConsiderHoursPerRestDay": "RWH",
    "BufferLateEarlyMinutes": "BL",
    "BufferDaysAllowedPerMonth": "BLD",
    "AttendanceRules": "ARule",
    "BreakDuringOvertimeMinutes": "OTB",
    "CalcProfessionalTax": "CalPT",
    "CalcProvidentFund": "CalPF",
    "CalcESIC": "CalESIC",
    "CalcTDS": "CalTDS",
    "IncomeTaxSlab": "SlabTDS",
    "LeavingTimeRounding": "LTimeROff",
    "SwipingScanningForMealBreak": "ScanMB",
    "fkSalAcctId": "fkSAcctId",
    "fkLoanAcctId": "fkLAcctId",
    "fkIncentiveAcctId": "fkIAcctId",
    "fkNoticeAcctId": "fkRAcctId",
    "Remarks": "Remarks",
    "SandwichRuleForLeaves": "Sandwich",
    "fkReportTo1EmpId": "fkEmp1Id",
    "fkReportTo2EmpId": "fkEmp2Id",
    "OTIncludeAllowance": "AllowanceBasic",
    "OTIncludeTravelAllowance": "TABasic",
    "OTIncludeHousingAllowance": "HABasic",
    "OTIncludeDearnessAllowance": "DABasic",
    "OTIncludeIncentive": "IncentiveBasic",
    "OTIncludeEducationAllowance": "EABasic",
    "OTIncludeMedicalAllowance": "MABasic",
    "OTIncludeOtherAllowance": "OtherBasic",
    "PFIncludeAllowance": "PFA",
    "PFIncludeTravelAllowance": "PFTA",
    "PFIncludeHousingAllowance": "PFHA",
    "PFIncludeIncentive": "PFI",
    "PFIncludeEducationAllowance": "PFEA",
    "PFIncludeMedicalAllowance": "PFMA",
    "PFIncludeOtherAllowance": "PFOA",
    "GovtHolidaysPartOfAllowances": "GHA",
    "RestDaysPartOfAllowances": "RDA",
    "IncentiveOnlyIfOvertimeFulfilled": "IORF",
    "OtherOnlyIfOvertimePerformed": "OAOP",
    "NoticeRetentionAmount": "Retention",
    "SuppliedTo": "fkFContId",
    "ManpowerAgency": "fkTContId",
    "GrossSalary": "SalGross",
    "PenaltyPerAbsentDay": "AbPenalty",
    "WorkingHoursVariant": "Variant",
    "RestDay1Variant": "RDVariant",
    "AdjustmentExtraWorkingHour": "EWHour",
    "LastYearExtraWorkingHour": "LYEWHour",
    "SetPF": "SetPF",
    "Latitude": "Latitude",
    "Longitude": "Longitude",
    "Altitude": "Altitude",
    "TDSDeductionPercent": "TDSDeduct",
    "MonthlyDeduction": "MDeduction",
    "DeductionDescription": "DedDescription",
    "fkAllowanceDesId": "fkDesId",
}

# These fields are part of the real table but are not currently exposed
# by the React form.  They therefore remain at their DB/default value
# when a new structure is created.
SAL_STRUCTURE_DEFAULTS = {
    "fkUserId": None,
    # SQL Server marks these overtime type flags as NOT NULL.
    "TOTI": False,
    "TOTII": False,
    "OtherBasic": False,
    "EOT": False,
    "EWHour": False,
    "MABasic": False,
    "EABasic": False,
    "IncentiveBasic": False,
    "DABasic": False,
    "HABasic": False,
    "TABasic": False,
    "AllowanceBasic": False,
    "Variant": False,
    "PFA": False,
    "PFTA": False,
    "PFHA": False,
    "PFI": False,
    "PFEA": False,
    "PFMA": False,
    "PFOA": False,
    "RDVariant": False,
    "SetPF": False,
    "Sandwich": False,
    "GHA": False,
    "RDA": False,
    "IORF": False,
    "OAOP": False,
}


def _salary_structure_type_to_db(value):
    """Convert the UI allowance-type labels to the legacy 5-char codes."""
    if value in (None, ""):
        return None
    text = str(value).strip()
    if text.lower() in {"fixed", "fixed amount"}:
        return "Fixed"
    if text.lower() in {"basic", "% of basic", "percentage", "percent"}:
        return "Basic"
    return text[:5]


def _salary_structure_type_from_db(value):
    if value in (None, ""):
        return ""
    text = str(value).strip()
    if text.lower() == "fixed":
        return "Fixed Amount"
    if text.lower() == "basic":
        return "% of Basic"
    return text


def _salary_structure_payload_to_db(payload: dict, current_user: dict, db: Session, *, creating: bool):
    """Translate API-facing Salary Structure fields to SalStructureTest columns."""
    source = dict(payload)
    result = {}

    for api_name, db_name in SAL_STRUCTURE_FIELD_MAP.items():
        if api_name not in source:
            continue
        value = source[api_name]
        if api_name in SAL_STRUCTURE_TYPE_FIELDS:
            value = _salary_structure_type_to_db(value)
        result[db_name] = value

    # The legacy table's audit user is server-controlled.
    result["fkUserId"] = resolve_legacy_user_id(db, current_user["pkid"])

    # These SQL Server columns are NOT NULL.  A blank checkbox/value from
    # the React form must therefore be stored as False rather than NULL.
    for flag_column in ("TOTI", "TOTII"):
        if flag_column in result and result[flag_column] in (None, ""):
            result[flag_column] = False

    # SalStructureTest.Revise is also NOT NULL.  The form does not need to
    # supply this audit/revision timestamp, so the backend owns it.  Use the
    # current server timestamp whenever the field is omitted or blank.
    if result.get("Revise") in (None, ""):
        result["Revise"] = datetime.now()

    if creating:
        for key, default in SAL_STRUCTURE_DEFAULTS.items():
            if key not in result and key != "fkUserId":
                result[key] = default

        # SalStructureTest requires every T* allowance type column to be
        # non-null.  The React form allows the user to leave the type blank,
        # so normalize an omitted/blank type to the legacy Fixed code instead
        # of sending NULL to SQL Server.  Do this for all allowance pairs so
        # partially filled new structures are still insertable.
        for type_column in (
            "TAllowance",
            "TTravelling",
            "THousing",
            "TDaily",
            "TIncentive",
            "TEducation",
            "TMedical",
            "TOther",
        ):
            if result.get(type_column) in (None, ""):
                result[type_column] = "Fixed"

    return result


def _salary_structure_to_api(structure):
    """Serialize SalStructureTest using the names expected by the React UI."""
    raw = salary_structure_to_dict(structure)
    if raw is None:
        return None

    data = {"pkSalStructureId": raw.get("pkSSId")}

    reverse_map = {db_name: api_name for api_name, db_name in SAL_STRUCTURE_FIELD_MAP.items()}
    for db_name, value in raw.items():
        api_name = reverse_map.get(db_name)
        if api_name:
            data[api_name] = value

    for api_name, db_name in SAL_STRUCTURE_TYPE_FIELDS.items():
        data[api_name] = _salary_structure_type_from_db(raw.get(db_name))

    # Explicitly expose fields that are not part of the UI/API mapping
    # only when useful for debugging/integration; this does not change
    # the existing UI contract.
    data["fkUserId"] = raw.get("fkUserId")
    data["SalStart"] = raw.get("SalStart")
    data["SalGross"] = raw.get("SalGross")

    return data


def _next_salary_structure_id(db: Session) -> int:
    """SalStructureTest.pkSSId is numeric, not an IDENTITY column."""
    latest = (
        db.query(SalaryStructure.pkSSId)
        .order_by(SalaryStructure.pkSSId.desc())
        .first()
    )
    if not latest or latest[0] is None:
        return 1
    return int(latest[0]) + 1


# --------------------------------------------------
# LIST SALARY STRUCTURES
# --------------------------------------------------

@router.get(
    "/salary-structures",
)
def list_salary_structures(
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    structures = get_salary_structures(db)

    return {
        "total": len(structures),
        "structures": [
            _salary_structure_to_api(structure)
            for structure in structures
        ],
    }


# --------------------------------------------------
# GET SALARY STRUCTURE
# --------------------------------------------------

@router.get(
    "/salary-structures/{structure_id}",
)
def get_salary_structure(
    structure_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    structure = get_salary_structure_by_id(db, structure_id)

    if not structure:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Salary structure not found",
        )

    return {
        "structure": _salary_structure_to_api(structure),
    }


# --------------------------------------------------
# CREATE SALARY STRUCTURE
# REQUIRES: salary_structure / add
# --------------------------------------------------

@router.post(
    "/salary-structures",
    status_code=status.HTTP_201_CREATED,
)
def create_new_salary_structure(
    request: SalStructureCreateRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_structure", "add")
    ),
):
    # The employee is the parent record this structure belongs to.
    if not get_salary_employee_by_id(db, request.fkEmpId):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Employee not found",
        )

    payload = request.model_dump(exclude_unset=True)
    payload.pop("paid_holidays", None)

    db_payload = _salary_structure_payload_to_db(
        payload,
        current_user,
        db,
        creating=True,
    )
    db_payload["pkSSId"] = _next_salary_structure_id(db)

    # Prevent a second structure with the same generated key in the
    # unlikely event another transaction created one between reads.
    if get_salary_structure_by_id(db, db_payload["pkSSId"]):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Could not allocate a new Salary Structure ID. Please retry.",
        )

    try:
        structure = create_salary_structure(db, db_payload)
    except (DataError, IntegrityError) as exc:
        print("=== SalStructureTest INSERT FAILED ===")
        print(repr(exc.orig) if hasattr(exc, "orig") else repr(exc))
        traceback.print_exc()
        print("========================================")
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "Salary structure could not be saved because one or more "
                "values do not match the SalStructureTest schema."
            ),
        ) from exc

    return {
        "message": "Salary structure created successfully",
        "structure": _salary_structure_to_api(structure),
    }


# --------------------------------------------------
# UPDATE SALARY STRUCTURE
# REQUIRES: salary_structure / edit
# --------------------------------------------------

@router.put(
    "/salary-structures/{structure_id}",
)
def update_existing_salary_structure(
    structure_id: int,
    request: SalStructureUpdateRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_structure", "edit")
    ),
):
    structure = get_salary_structure_by_id(db, structure_id)

    if not structure:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Salary structure not found",
        )

    update_payload = request.model_dump(exclude_unset=True)
    update_payload.pop("paid_holidays", None)

    if "fkEmpId" in update_payload:
        if update_payload["fkEmpId"] is None:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Employee is required",
            )
        if not get_salary_employee_by_id(db, update_payload["fkEmpId"]):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Employee not found",
            )

    db_payload = _salary_structure_payload_to_db(
        update_payload,
        current_user,
        db,
        creating=False,
    )

    try:
        structure = update_salary_structure(
            db,
            structure,
            **db_payload,
        )
    except (DataError, IntegrityError) as exc:
        print("=== SalStructureTest UPDATE FAILED ===")
        print(repr(exc.orig) if hasattr(exc, "orig") else repr(exc))
        traceback.print_exc()
        print("========================================")
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "Salary structure could not be updated because one or more "
                "values do not match the SalStructureTest schema."
            ),
        ) from exc

    return {
        "message": "Salary structure updated successfully",
        "structure": _salary_structure_to_api(structure),
    }


# --------------------------------------------------
# DELETE SALARY STRUCTURE
# REQUIRES: salary_structure / delete
# --------------------------------------------------

@router.delete(
    "/salary-structures/{structure_id}",
)
def delete_existing_salary_structure(
    structure_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_permission("salary_structure", "delete")
    ),
):
    structure = get_salary_structure_by_id(db, structure_id)

    if not structure:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Salary structure not found",
        )

    try:
        delete_salary_structure(db, structure)
    except (DataError, IntegrityError) as exc:
        print("=== SalStructureTest DELETE FAILED ===")
        print(repr(exc.orig) if hasattr(exc, "orig") else repr(exc))
        traceback.print_exc()
        print("========================================")
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Salary structure could not be deleted.",
        ) from exc

    return {
        "message": "Salary structure deleted successfully",
    }


# ==================================================
# EMPLOYEE CONTACTS
# ==================================================

@router.get(
    "/salary-employees/{emp_id}/contacts",
)
def list_employee_contacts(

    emp_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    # NOTE: read access follows the same app-wide pattern
    # as every other GET — see the pending view-permission
    # TODO in security.py.

    if not get_salary_employee_by_id(db, emp_id):

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Employee not found",
        )

    return {

        "success": True,

        "contacts": [

            employee_contact_to_dict(row)

            for row in get_employee_contacts(db, emp_id)
        ],
    }


@router.put(
    "/salary-employees/{emp_id}/contacts",
)
def replace_employee_contact_rows(

    emp_id: int,

    request: SalEmpContactListRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("salary_employee", "edit")
    ),
):

    if not get_salary_employee_by_id(db, emp_id):

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Employee not found",
        )

    try:

        rows = replace_employee_contacts(
            db,
            emp_id,
            normalize_contact_rows(request.rows),
        )

    except (DataError, IntegrityError) as exc:

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "Contact rows could not be saved because one or "
                "more values do not match the SalEmpContact "
                "table schema."
            ),
        ) from exc

    return {

        "message": "Contacts saved successfully",

        "contacts": [
            employee_contact_to_dict(row)
            for row in rows
        ],
    }


# ==================================================
# EMPLOYEE RELATIVES
# ==================================================

@router.get(
    "/salary-employees/{emp_id}/relations",
)
def list_salemp_relations(

    emp_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    if not get_salary_employee_by_id(db, emp_id):

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Employee not found",
        )

    return {

        "success": True,

        "relations": [

            salemp_relation_to_dict(row)

            for row in get_salemp_relations(db, emp_id)
        ],
    }


@router.put(
    "/salary-employees/{emp_id}/relations",
)
def replace_salemp_relation_rows(

    emp_id: int,

    request: SalEmpRelationListRequest,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("salary_employee", "edit")
    ),
):

    if not get_salary_employee_by_id(db, emp_id):

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Employee not found",
        )

    try:

        rows = replace_employee_relations(
            db,
            emp_id,
            normalize_relation_rows(request.rows),
        )

    except (DataError, IntegrityError) as exc:

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "Relative rows could not be saved because one "
                "or more values do not match the SalEmpRelation "
                "table schema."
            ),
        ) from exc

    return {

        "message": "Relatives saved successfully",

        "relations": [
            salemp_relation_to_dict(row)
            for row in rows
        ],
    }


# ==================================================
# EMPLOYEE DOCUMENTS
# ==================================================

@router.get(
    "/salary-employees/{emp_id}/documents",
)
def list_employee_documents(

    emp_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    if not get_salary_employee_by_id(db, emp_id):

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Employee not found",
        )

    return {

        "success": True,

        "documents": [

            employee_document_payload(emp_id, row)

            for row in get_employee_documents(db, emp_id)
        ],
    }


@router.post(
    "/salary-employees/{emp_id}/documents",
    status_code=status.HTTP_201_CREATED,
)
def upload_employee_document(

    emp_id: int,

    file: UploadFile = File(...),

    fkDTId: int | None = Form(default=None),

    ValidUntil: str | None = Form(default=None),

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("salary_employee", "edit")
    ),
):

    if not get_salary_employee_by_id(db, emp_id):

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Employee not found",
        )


    # ==============================================
    # VALIDATE EXTENSION
    # ==============================================

    suffix = Path(
        file.filename or ""
    ).suffix.lower()

    if suffix not in ALLOWED_DOCUMENT_EXTENSIONS:

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "Unsupported file type. Allowed: "
                + ", ".join(
                    sorted(ALLOWED_DOCUMENT_EXTENSIONS)
                )
            ),
        )


    # ==============================================
    # PARSE ValidUntil
    # ==============================================

    valid_until = None

    if ValidUntil:

        try:

            valid_until = datetime.fromisoformat(
                ValidUntil
            )

        except ValueError:

            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=(
                    "Valid Until must be a date in "
                    "YYYY-MM-DD format."
                ),
            )


    # ==============================================
    # WRITE FILE
    # ==============================================

    stored_name = _safe_document_name(
        file.filename
    )

    target_dir = _employee_document_dir(emp_id)

    target_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    target_path = target_dir / stored_name

    written = 0

    try:

        with target_path.open("wb") as handle:

            while True:

                chunk = file.file.read(
                    1024 * 1024
                )

                if not chunk:

                    break

                written += len(chunk)

                if written > MAX_DOCUMENT_BYTES:

                    handle.close()

                    target_path.unlink(
                        missing_ok=True
                    )

                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail=(
                            "File is too large. Maximum size "
                            "is 10 MB."
                        ),
                    )

                handle.write(chunk)

    except HTTPException:

        raise

    except OSError as exc:

        target_path.unlink(
            missing_ok=True
        )

        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Could not store the uploaded file.",
        ) from exc

    finally:

        file.file.close()


    # ==============================================
    # SAVE ROW
    #
    # If the insert fails, remove the orphaned file so
    # disk and table stay in step.
    # ==============================================

    try:

        document = create_employee_document(
            db,
            emp_id,
            {
                "fkDTId": fkDTId,
                "DocFile": stored_name,
                "ValidUntil": valid_until,
            },
        )

    except (DataError, IntegrityError) as exc:

        target_path.unlink(
            missing_ok=True
        )

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "Document could not be saved because one or "
                "more values do not match the SalEmpDocuments "
                "table schema."
            ),
        ) from exc

    return {

        "message": "Document uploaded successfully",

        "document": employee_document_payload(
            emp_id,
            document,
        ),
    }


@router.get(
    "/salary-employees/{emp_id}/documents/{document_id}/file",
)
def download_employee_document(

    emp_id: int,

    document_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        get_current_user
    ),
):

    document = get_employee_document(
        db,
        emp_id,
        document_id,
    )

    if not document:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Document not found",
        )

    stored_name = document.DocFile or ""

    file_path = (
        _employee_document_dir(emp_id) / stored_name
    )


    # ==============================================
    # PATH SAFETY
    #
    # DocFile comes from the DB, and legacy rows may hold
    # anything at all — make sure the resolved path is
    # still inside this employee's folder before serving.
    # ==============================================

    try:

        resolved = file_path.resolve()

        allowed_root = _employee_document_dir(
            emp_id
        ).resolve()

        inside = (
            resolved == allowed_root
            or allowed_root in resolved.parents
        )

    except OSError:

        inside = False

    if not stored_name or not inside or not resolved.is_file():

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=(
                "The file for this document is missing "
                "from storage."
            ),
        )

    media_type, _ = mimetypes.guess_type(
        str(resolved)
    )

    display_name = document_display_name(
        stored_name
    )

    # inline so PDFs and images open in a browser tab
    # rather than forcing a download.

    return FileResponse(
        path=resolved,
        media_type=media_type or "application/octet-stream",
        filename=display_name,
        content_disposition_type="inline",
    )


@router.delete(
    "/salary-employees/{emp_id}/documents/{document_id}",
)
def delete_employee_document_row(

    emp_id: int,

    document_id: int,

    db: Session = Depends(get_db),

    current_user: dict = Depends(
        require_permission("salary_employee", "edit")
    ),
):

    document = get_employee_document(
        db,
        emp_id,
        document_id,
    )

    if not document:

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Document not found",
        )

    stored_name = document.DocFile or ""

    delete_employee_document(
        db,
        document,
    )

    if stored_name:

        (
            _employee_document_dir(emp_id) / stored_name
        ).unlink(
            missing_ok=True
        )

    return {

        "message": "Document deleted successfully",
    }